#!/usr/bin/env python3
"""Valida estrutura e conteúdo mínimo dos outputs Office do Report Live."""

from __future__ import annotations

import argparse
import json
import math
import zipfile
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader


def assert_numeric_series(actual: list[float], expected: list[float], owner: str) -> None:
    if len(actual) != len(expected):
        raise ValueError(
            f"Gráfico {owner} com {len(actual)} pontos; esperado {len(expected)}"
        )
    for index, (actual_value, expected_value) in enumerate(zip(actual, expected), start=1):
        if not math.isclose(float(actual_value), float(expected_value), rel_tol=1e-9, abs_tol=1e-7):
            raise ValueError(
                f"Gráfico {owner} diverge no ponto {index}: {actual_value} != {expected_value}"
            )


def validate(plan_path: Path, pptx_path: Path, pdf_path: Path) -> dict[str, object]:
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    slide_ids = tuple(item["slide_instance_id"] for item in plan["slides"])
    if not slide_ids or len(set(slide_ids)) != len(slide_ids):
        raise ValueError(f"Slice vazio ou duplicado: {slide_ids}")

    if not zipfile.is_zipfile(pptx_path):
        raise ValueError("PPTX inválido: pacote ZIP ausente")
    prs = Presentation(str(pptx_path))
    if len(prs.slides) != len(slide_ids):
        raise ValueError(f"PPTX com {len(prs.slides)} slides; esperado {len(slide_ids)}")
    charts = sum(
        1 for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_chart", False)
    )
    expected_charts = sum(slide_id in {"c4", "m1", "m2"} for slide_id in slide_ids)
    if charts != expected_charts:
        raise ValueError(f"PPTX com {charts} gráficos nativos; esperado {expected_charts}")

    by_id = {item["slide_instance_id"]: item for item in plan["slides"]}
    for slide_id, slide in zip(slide_ids, prs.slides):
        chart_shapes = [shape for shape in slide.shapes if getattr(shape, "has_chart", False)]
        if slide_id not in {"c4", "m1", "m2"}:
            if chart_shapes:
                raise ValueError(f"Slide {slide_id} contém gráfico nativo inesperado")
            continue
        if len(chart_shapes) != 1:
            raise ValueError(f"Slide {slide_id} contém {len(chart_shapes)} gráficos; esperado 1")
        actual_series = [list(series.values) for series in chart_shapes[0].chart.series]
        if slide_id == "c4":
            expected_series = [series["values"] for series in by_id[slide_id]["chart"]["series"]]
        else:
            expected_series = [[row["value"] for row in by_id[slide_id]["ranking"]]]
        if len(actual_series) != len(expected_series):
            raise ValueError(
                f"Gráfico {slide_id} com {len(actual_series)} séries; esperado {len(expected_series)}"
            )
        for actual, expected in zip(actual_series, expected_series):
            assert_numeric_series(actual, expected, slide_id)
    pptx_text = "\n".join(
        shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    ).lower()

    reader = PdfReader(str(pdf_path))
    if len(reader.pages) != len(slide_ids):
        raise ValueError(f"PDF com {len(reader.pages)} páginas; esperado {len(slide_ids)}")
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages).lower()
    expected_text = ["report live" if item["slide_instance_id"] == "c0" else item["title"].lower()
                     for item in plan["slides"]]
    for expected in expected_text:
        if expected not in pptx_text:
            raise ValueError(f"Texto ausente no PPTX: {expected}")
        if expected not in pdf_text:
            raise ValueError(f"Texto ausente no PDF: {expected}")

    page_sizes = [
        (round(float(page.mediabox.width)), round(float(page.mediabox.height)))
        for page in reader.pages
    ]
    if any(size != (960, 540) for size in page_sizes):
        raise ValueError(f"Canvas PDF inesperado: {page_sizes}")

    return {
        "ok": True,
        "run_id": plan["run_id"],
        "slides": list(slide_ids),
        "pptx_slides": len(prs.slides),
        "native_charts": charts,
        "pdf_pages": len(reader.pages),
        "pdf_canvas_points": [960, 540],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(validate(args.plan, args.pptx, args.pdf), ensure_ascii=False))


if __name__ == "__main__":
    main()
