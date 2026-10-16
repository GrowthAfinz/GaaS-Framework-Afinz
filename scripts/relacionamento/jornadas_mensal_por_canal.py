"""Resultado mensal por canal das jornadas de relacionamento (Desbloqueio, Incentivo ao Uso,
Welcome, Pool de Ofertas) a partir do export do BI Framework_Growth_Aquisicao_CartaoDinamica.

Uso: python scripts/relacionamento/jornadas_mensal_por_canal.py <export.csv> [saida.xlsx]

O export traz 5 tabelas dinâmicas lado a lado (WPP, E-mail, SMS, Resultados, Push). O pivô
replica a mesma activity em todas as jornadas que a usam, então o script só lê as jornadas
oficiais (ver TABS) e deduplica por activity + data.
"""
import sys
import csv,sys,pandas as pd
F=sys.argv[1]
rows=list(csv.reader(open(F,encoding='latin-1'),delimiter=';'))
def g(r,j): return r[j].strip() if j<len(r) else ''
def block(start_row, cols, names):
    out=[]
    for r in rows[start_row:]:
        vals=[g(r,j) for j in cols]
        if not any(vals): continue
        out.append(vals)
    return pd.DataFrame(out,columns=names)
B={
 'wpp':  block(4,[0,1,2,3,4,5],['journey','activity','date','sent','delivered','read']),
 'email':block(3,[8,9,10,11,12,13,14],['journey','activity','date','sent','delivered','opens','clicks']),
 'sms':  block(3,[16,17,18,19,20,21],['journey','activity','date','sent','delivered','undelivered']),
 'res':  block(1,[23,24,25,26,27,28,29,30,31],['journey','activity','date','canal','propostas','aprovados','finalizados','assistida','independente']),
 'push': block(3,[35,36,37,38,39,40,41,42],['journey','activity','date','attempts','success','fail','notfound','unregistered']),
}
for k,df in B.items():
    df['dt']=pd.to_datetime(df['date'],format='%d/%m/%Y',errors='coerce')
    df['mes']=df['dt'].dt.to_period('M')
    for c in df.columns:
        if c not in('journey','activity','date','canal','dt','mes'):
            df[c]=pd.to_numeric(df[c].str.replace('.','',regex=False).str.replace(',','.',regex=False),errors='coerce').fillna(0)
import re
VC='B2C/B2B2C (VC)'; PLU='Plurix (Mais Amigo)'
TABS=[
 ('Desbloqueio',[
   (VC,'Jornada de Desbloqueio VC',['JOR_ATIVACAO_DESBLOQUEIO_VC'],''),
   (PLU,'Jornada de Desbloqueio Plurix',['JOR_ATIVAÇÃO_DESBLOQUEIO_PLURIX_MAISAMIGO'],''),
   (PLU,'Pós-Tombamento Plurix (separado)',['JOR_POS_TOMBAMENTO_DESBLOQUEIO_PLURIX_MAISAMIGO','JOR_POS_TOMBAMENTO_DESBLOQUEIO_PLURIX_MAISAMIGO_MAIO'],'Ação pontual de desbloqueio da base migrada (tombamento): nov/25–jan/26 e mai–jul/26. Fora da régua perene.'),
 ]),
 ('Incentivo ao Uso',[
   (VC,'Jornada de Incentivo ao Uso VC',['JOR_ATIVACAO_INCENTIVO_AO_USO_AFINZ_VC'],'Jornada perene só tem e-mail, com disparos de mar/26 a mai/26 e nenhum depois disso no arquivo.'),
   (VC,'Campanhas Copa · Ativação S1, S2 + Reativação S1 (separado)',['JOR_RENTABILIZACAO_ATIVACAO_B2C_B2B2C_COPA_SEM1','JOR_RENTABILIZACAO_ATIVACAO_B2C_B2B2C_COPA_SEM2','JOR_RENTABILIZACAO_REATIVACAO_B2C_B2B2C_COPA_SEM1'],'Campanhas sazonais, não jornada perene. Excluída a activity afz_car_vis_rnt_wpp_nov_disp2s1copa_diarioo: é da jornada NOVOS e aparecia replicada na Reativação.'),
   (PLU,'Jornada de Incentivo ao Uso Plurix',['JOR_INCENTIVO_AO_USO_PLURIX_MAISAMIGO'],''),
 ]),
 ('Welcome',[
   (VC,'Jornada Welcome VC',['JOR_ATIVACAO_WELCOME_AFINZ_VC'],''),
   (VC,'Errata Welcome VC (separado)',['ERRATA_JOR_ATIVACAO_WELCOME_AFINZ_VC'],'Disparo pontual de correção para a base ativa (jun–jul/26). Não faz parte da régua.'),
   (PLU,'Jornada Welcome Plurix',['JOR_ATIVACAO_WELCOME_PLURIX_MAISAMIGO'],'Inclui as activities renomeadas com sufixo _diariio/_diiario, ativas a partir de ago/26.'),
 ]),
 ('Pool de Ofertas',[
   (VC,'Pool de Ofertas · Shopping Afinz (e-mail)',['JOR_RENT_POOL_OFERTAS'],''),
   (VC,'Repique Emagrecentro',['JOR_RENT_POOL_OFERTAS_REPIQUE_EMAGRECENTRO'],''),
   (VC,'Shopping Afinz 2025 · Emagrecentro',['JOR_SHOPPING_AFINZ_EMAGRECENTRO'],'Versão anterior do pool (out–nov/25).'),
   (VC,'Shopping Afinz 2025 · Star Móveis',['JOR_SHOPPING_AFINZ_STARMOVEIS'],'Versão anterior do pool (out–nov/25).'),
   (VC,'Shopping Afinz 2025 · Virtuosa',['JOR_SHOPPING_AFINZ_VIRTUOSA'],'Versão anterior do pool (out–nov/25).'),
   (PLU,'Pool de Ofertas Plurix · Ofertas da Semana',['JOR_CIC_POOL_OFERTAS_PLURIX'],''),
   (PLU,'POOL MANUAL (provável teste)',['POOL MANUAL'],'Apenas 16 envios em 2 dias (22–23/09/26), com a mesma activity da CIC Plurix logo antes da estreia dela. Tem cara de teste: avaliar antes de considerar.'),
 ]),
]
EXCL_ACTS={'afz_car_vis_rnt_wpp_nov_disp2s1copa_diarioo'}
def etapa(act):
    a=act.lower()
    m=re.search(r'(?:^|[_\-])d(\d{1,2})(?=[\-_a-z+:])',a) or re.search(r'(?<=[a-z])d(\d{1,2})(?=[a-z+\-_:])',a)
    return int(m.group(1)) if m else None
def etapa_label(act):
    a=act.lower()
    n=etapa(act)
    if n is None:
        m=re.search(r'd(\d{1,2})_\d{8}$',a)
        if m: n=int(m.group(1))
    if n is not None: return f'D{n}',(0,n)
    for pat,lab in [(r'clique(\d)','Clique {}'),(r'repique\w*?(\d)_','Repique {}'),(r'disp(\d)','Disparo {}'),(r'postombamento(\d)','Disparo {}'),(r'disparo(\d)','Disparo {}')]:
        m=re.search(pat,a)
        if m: return lab.format(m.group(1)),(1,int(m.group(1)))
    if 'postombamento_pontual' in a: return 'Disparo 1',(1,1)
    return 'Único',(2,0)

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

MESES=['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']
CH=[ # key, label, [(col,label)], [(label, num, den)]
 ('wpp','WHATSAPP',[('sent','Enviados'),('delivered','Entregues'),('read','Lidos')],[('% Entrega','Entregues','Enviados'),('% Leitura','Lidos','Entregues')]),
 ('email','E-MAIL',[('sent','Enviados'),('delivered','Entregues'),('opens','Aberturas únicas'),('clicks','Cliques únicos')],[('% Entrega','Entregues','Enviados'),('% Abertura','Aberturas únicas','Entregues'),('CTR','Cliques únicos','Entregues'),('CTOR','Cliques únicos','Aberturas únicas')]),
 ('sms','SMS',[('sent','Enviados'),('delivered','Entregues'),('undelivered','Não entregues')],[('% Entrega','Entregues','Enviados'),('% Falha','Não entregues','Enviados')]),
 ('push','PUSH',[('attempts','Enviados'),('success','Sucesso'),('fail','Falha'),('unregistered','Não registrados')],[('% Sucesso','Sucesso','Enviados'),('% Falha','Falha','Enviados'),('% Não registrados','Não registrados','Enviados')]),
]
# clean data
D={}
for k,df in B.items():
    if k=='res': continue
    df=df[df.dt.notna()&(df.journey!='Total Geral')&~df.activity.isin(EXCL_ACTS)].copy()
    D[k]=df

F='Arial'
f_title=Font(name=F,size=16,bold=True,color='0F172A')
f_sub=Font(name=F,size=9,italic=True,color='475569')
f_bu=Font(name=F,size=12,bold=True,color='FFFFFF'); fill_vc=PatternFill('solid',fgColor='0E7490'); fill_plu=PatternFill('solid',fgColor='6D28D9')
f_sec=Font(name=F,size=11,bold=True,color='0F172A')
f_note=Font(name=F,size=9,italic=True,color='92400E')
f_ch=Font(name=F,size=10,bold=True,color='0F172A'); fill_ch=PatternFill('solid',fgColor='E2E8F0')
f_hd=Font(name=F,size=9,bold=True,color='FFFFFF'); fill_hd=PatternFill('solid',fgColor='334155')
f_cell=Font(name=F,size=10); f_tot=Font(name=F,size=10,bold=True); fill_tot=PatternFill('solid',fgColor='F1F5F9')
thin=Side(style='thin',color='CBD5E1'); top=Border(top=Side(style='medium',color='334155'))
NUM='#,##0;-#,##0;"-"'; PCT='0.0%;-0.0%;"-"'
ctr=Alignment(horizontal='center',vertical='center',wrap_text=True)

def mes_label(p): return f'{MESES[p.month-1]}/{str(p.year)[2:]}'
def put(ws,r,c,v,font=f_cell,fmt=None,fill=None,al=None,border=None):
    x=ws.cell(r,c,v); x.font=font
    if fmt: x.number_format=fmt
    if fill: x.fill=fill
    if al: x.alignment=al
    if border: x.border=border
    return x

def header(ws,r,labels):
    for i,l in enumerate(labels): put(ws,r,1+i,l,f_hd,fill=fill_hd,al=ctr)
    ws.row_dimensions[r].height=28

def rate_formula(r,cols,num,den):
    n=L(cols[num]); d=L(cols[den]); return f'=IFERROR({n}{r}/{d}{r},0)'

def monthly_table(ws,r,df,ch):
    key,lab,mets,rates=ch
    put(ws,r,1,f'▸ {lab}',f_ch,fill=fill_ch)
    for c in range(2,2+len(mets)+len(rates)): ws.cell(r,c).fill=fill_ch
    r+=1
    labels=['Mês']+[m[1] for m in mets]+[x[0] for x in rates]
    header(ws,r,labels); r+=1
    cols={l:i+1 for i,l in enumerate(labels)}
    g=df.groupby('mes')[[m[0] for m in mets]].sum()
    months=pd.period_range(g.index.min(),g.index.max(),freq='M')
    g=g.reindex(months,fill_value=0)
    first=r
    for p,row in g.iterrows():
        put(ws,r,1,mes_label(p),al=Alignment(horizontal='left'))
        for i,(c,_) in enumerate(mets): put(ws,r,2+i,int(row[c]),fmt=NUM)
        for j,(rl,nu,de) in enumerate(rates): put(ws,r,2+len(mets)+j,rate_formula(r,cols,nu,de),fmt=PCT)
        r+=1
    last=r-1
    put(ws,r,1,'TOTAL',f_tot,fill=fill_tot,border=top)
    for i in range(len(mets)):
        cl=L(2+i); put(ws,r,2+i,f'=SUM({cl}{first}:{cl}{last})',f_tot,NUM,fill_tot,border=top)
    for j,(rl,nu,de) in enumerate(rates): put(ws,r,2+len(mets)+j,rate_formula(r,cols,nu,de),f_tot,PCT,fill_tot,border=top)
    return r+2, int(g[mets[0][0]].sum())

def etapa_table(ws,r,rows,ch):
    key,lab,mets,rates=ch
    put(ws,r,1,f'▸ {lab} · por etapa (acumulado do período)',f_ch,fill=fill_ch)
    ncol=13
    for c in range(2,ncol+1): ws.cell(r,c).fill=fill_ch
    r+=1
    labels=['Etapa']+[m[1] for m in mets]+[x[0] for x in rates]
    header(ws,r,labels)
    for i,l in enumerate(['Início','Fim','Jornada / seção','Activity']): put(ws,r,10+i,l,f_hd,fill=fill_hd,al=ctr)
    for c in range(len(labels)+1,10): ws.cell(r,c).fill=fill_hd
    r+=1
    cols={l:i+1 for i,l in enumerate(labels)}
    for e in rows:
        put(ws,r,1,e['etapa'],al=Alignment(horizontal='left'))
        for i,(c,_) in enumerate(mets): put(ws,r,2+i,int(e[c]),fmt=NUM)
        for j,(rl,nu,de) in enumerate(rates): put(ws,r,2+len(mets)+j,rate_formula(r,cols,nu,de),fmt=PCT)
        b=10
        put(ws,r,b,e['ini']); put(ws,r,b+1,e['fim']); put(ws,r,b+2,e['sec']); put(ws,r,b+3,e['act'])
        r+=1
    return r+1

wb=Workbook(); wsM=wb.active; wsM.title='Mapa & Premissas'
mapa=[]
for tab,secs in TABS:
    ws=wb.create_sheet(tab)
    ws.sheet_view.showGridLines=False
    put(ws,1,1,f'{tab} · Resultado mensal por canal',f_title)
    put(ws,2,1,'Fonte: Framework_Growth_Aquisição_CartaoDinamica_BI_4.csv (blocos de canal WPP, E-mail, SMS e Push). Linhas deduplicadas por activity + data. Taxas calculadas sobre os volumes somados. Enviados = impactos, não clientes únicos. Regras completas na aba Mapa & Premissas.',f_sub)
    r=4; cur_bu=None; etapas={}
    for bu,title,js,note in secs:
        if bu!=cur_bu:
            put(ws,r,1,bu.upper(),f_bu,fill=fill_vc if bu==VC else fill_plu)
            for c in range(2,14): ws.cell(r,c).fill=fill_vc if bu==VC else fill_plu
            r+=2; cur_bu=bu
        put(ws,r,1,f'■ {title}',f_sec); r+=1
        put(ws,r,1,'Jornada(s): '+' · '.join(js),f_sub); r+=1
        if note: put(ws,r,1,'Nota: '+note,f_note); r+=1
        r+=1
        for ch in CH:
            key=ch[0]
            df=D[key][D[key].journey.isin(js)].drop_duplicates(['activity','date'])
            if df.empty: continue
            r,tot=monthly_table(ws,r,df,ch)
            mapa.append([tab,bu,title,' · '.join(js),ch[1],mes_label(df.mes.min()),mes_label(df.mes.max()),df.activity.nunique(),tot])
            for act,a in df.groupby('activity'):
                lab,sk=etapa_label(act)
                e={'etapa':lab,'sk':(bu!=VC,secs.index((bu,title,js,note)),sk,act),'act':act,'sec':title,
                   'ini':a.dt.min().strftime('%d/%m/%y'),'fim':a.dt.max().strftime('%d/%m/%y')}
                for c,_ in ch[2]: e[c]=a[c].sum()
                etapas.setdefault(key,[]).append(e)
        r+=1
    # etapa section
    put(ws,r,1,'POR ETAPA DA RÉGUA · acumulado do período inteiro',f_bu,fill=PatternFill('solid',fgColor='0F172A'))
    for c in range(2,14): ws.cell(r,c).fill=PatternFill('solid',fgColor='0F172A')
    r+=1
    put(ws,r,1,'Etapa extraída do nome da activity (dN = dias desde a entrada na jornada). Campanhas sem "dN" usam o número do disparo. Activities renomeadas aparecem em linhas próprias.',f_sub); r+=2
    for ch in CH:
        if ch[0] in etapas:
            rows=sorted(etapas[ch[0]],key=lambda e:e['sk'])
            r=etapa_table(ws,r,rows,ch)
    ws.column_dimensions['A'].width=16
    for c in range(2,10): ws.column_dimensions[L(c)].width=14
    for c,w in zip('JKLM',[10,10,48,78]): ws.column_dimensions[c].width=w
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True

# ---------- Mapa & Premissas ----------
ws=wsM; ws.sheet_view.showGridLines=False
put(ws,1,1,'Mapa de Jornadas & Premissas',f_title)
put(ws,2,1,'Resultado mensal por canal das frentes de Desbloqueio, Incentivo ao Uso, Welcome e Pool de Ofertas. Gerado em 16/10/2026.',f_sub)
r=4
put(ws,r,1,'PREMISSAS',f_sec); r+=1
prem=[
 'Fonte: Framework_Growth_Aquisição_CartaoDinamica_BI_4.csv. São 5 tabelas dinâmicas lado a lado: WhatsApp, E-mail, SMS, Resultados de performance e Push. As linhas "Total Geral" foram descartadas.',
 'Deduplicação: no export, a mesma activity aparece replicada em todas as jornadas que a usam (ex.: Desbloqueio Plurix = _SEEDLIST nas 1.471 linhas). Por isso só entra a jornada oficial, e cada tabela é deduplicada por activity + data.',
 'O bloco de Resultados (propostas/aprovados/finalizados) cobre só aquisição; nenhuma destas jornadas aparece nele. O output é de engajamento por canal.',
 'Enviados = impactos (mensagens), não clientes únicos. Um cliente que recebe D0, D5 e D15 conta 3 vezes.',
 'Taxas sobre volumes somados: % Entrega = Entregues ÷ Enviados · % Leitura = Lidos ÷ Entregues · % Abertura = Aberturas ÷ Entregues · CTR = Cliques ÷ Entregues · CTOR = Cliques ÷ Aberturas · Push: % Sucesso/Falha/Não registrados ÷ Enviados.',
 'Push: "Enviados" corresponde a total_attempts do export (padronizado com os demais canais).',
 'SMS e Push não trazem abertura nem clique no export. Só WhatsApp (leitura) e E-mail (abertura/clique) têm engajamento.',
 'Cada tabela começa no 1º mês com disparo daquela jornada naquele canal. Meses sem disparo no meio da série aparecem zerados (-).',
 'Disponibilidade por canal no export: WhatsApp set/25→out/26 · SMS set/25→out/26 · E-mail mar/26→out/26 · Push ago/26→out/26. Os pivôs estavam com filtro "Ano-Mês (Vários itens)", então meses fora do filtro não vieram.',
 'Em volumes muito pequenos, o export pode trazer Lidos maior que Entregues (ex.: Pós-Tombamento jan/26: 1 entregue e 7 lidos), o que gera % Leitura acima de 100%. É inconsistência da fonte e foi mantida como veio.',
 'Outubro/26 é parcial: dados até 15–16/10/2026.',
 '"VC" agrupa B2C e B2B2C (cartão Afinz Visa). O nome da jornada não permite separar as duas BUs.',
]
for p in prem:
    put(ws,r,1,'• '+p,f_cell); ws.cell(r,1).alignment=Alignment(wrap_text=True,vertical='top'); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=9); ws.row_dimensions[r].height=30; r+=1
r+=1
put(ws,r,1,'JORNADAS INCLUÍDAS',f_sec); r+=1
hd=['Aba','BU','Seção','Jornada(s)','Canal','Início','Fim','Nº activities','Enviados (total)']
header(ws,r,hd); r+=1
for m in mapa:
    for i,v in enumerate(m): put(ws,r,1+i,v,fmt=NUM if i>=7 else None)
    r+=1
r+=1
put(ws,r,1,'JORNADAS EXCLUÍDAS',f_sec); r+=1
header(ws,r,['Jornada','Motivo']); r+=1
exc=[
 ('JOR_ATIVACAO_DESBLOQUEIO_PLURIX_MAISAMIGO_SEEDLIST','Seedlist (QA). WPP é cópia idêntica da jornada oficial'),
 ('JOR_ATIVACAO_WELCOME_PLURIX_MAISAMIGO_SEEDLIST','Seedlist (QA). WPP é cópia idêntica da jornada oficial'),
 ('JOR_INCENTIVO_AO_USO_PLURIX_MAISAMIGO_TESTE / (Copiar)','Teste ou cópia. Linhas idênticas à jornada oficial'),
 ('JOR_Cartao_VC_Welcome / (Copiar)','Nome antigo ou cópia do Welcome VC. Mesma activity e mesmos números'),
 ('JOR_TESTE / JOR_TESTE_UX','Jornadas de teste. A UX replica a activity do Welcome VC'),
 ('JOR_ATIVACAO_PLURIX_COPA_TEST_* / teste_JOR_RENTABILIZACAO_*','Testes'),
 ('JOR_TRS_ATENDIMENTO_CENTRAL_DESBLOQUEAR','Transacional (resposta de atendimento), não é régua de desbloqueio'),
 ('JOR_TRS_POOL_OFERTAS_RELATORIO','E-mail de relatório interno do pool (107 envios)'),
 ('JOR_Upsell_Desbloqueio · JOR_VP_Cartao_Desbloqueio(_DaSanta/_Hirota) · JOR_Upsell_Ativacao · JOR_VP_Cartao_Ativacao','Pilotos Platinum de set–out/25, com menos de 600 envios no total'),
 ('afz_car_vis_rnt_wpp_nov_disp2s1copa_diarioo (activity)','Pertence às jornadas RENTABILIZACAO_NOVOS_COPA; aparecia replicada na Reativação Copa S1'),
]
for a,b in exc:
    put(ws,r,1,a); put(ws,r,2,b); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=9); r+=1
r+=1
put(ws,r,1,'RELACIONADAS, FORA DO ESCOPO (não solicitadas)',f_sec); r+=1
for a in ['JOR_RENTABILIZACAO_NOVOS_B2C_B2B2C_COPA_SEM1/SEM2','JOR_RENTABILIZACAO_ATIVACAO_*_FASE1_COPA_*, _BB_COPA_SEM1, _CADASTRO_INCOMPLETO_COPA, INCENTIVO_CADASTRO_COPA_PONTUAL','JOR_ATIVACAO_PLURIX_COPA_SEM1 (ativação Copa Plurix: WPP + e-mail, jun–ago/26)','JOR_*_REATIVACAO_* de aquisição (carrinho/abandono): são reativação de proposta, não de cartão']:
    put(ws,r,1,'• '+a); r+=1
for c,w in zip('ABCDEFGHI',[62,20,50,70,12,10,10,13,16]): ws.column_dimensions[c].width=w
for row in ws.iter_rows(min_row=4):
    for c in row:
        if c.column==1 and c.font.b is False: pass
ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
out=sys.argv[2] if len(sys.argv)>2 else 'outputs/Jornadas_Relacionamento_Mensal_por_Canal.xlsx'
wb.save(out); print(out)
