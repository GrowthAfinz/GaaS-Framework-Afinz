# -*- coding: utf-8 -*-
"""V2 TOPO DE FUNIL DINÂMICO: E-mails 3 e 4 a partir da aba "Briefing Fábrica" (linhas 9 a 46),
mais os ajustes de régua que vieram com eles (negrito nos quadros, saudação do E-mail 2, template).

Uso: python scripts/email-factory/build_topo_dinamico_v2_e3e4.py <planilha.xlsx> <v13_atual.html> <v13_novo.html> <saida.json> [saida.sql]

- Lê por rótulo: a coluna pelo cabeçalho "E-mail N" (linha 4) e o campo pela coluna "Onde fica".
- Negrito vira <b>, quebra vira <br>, PRI_NOME vira %%=v(@FirstName)=%% (mesmas regras da V1).
- Linhas automáticas (versões por rede, logo, faixa, pool, assinatura, nota legal, rodapé) não entram:
  links, UTM, assinatura, nota legal e logo de cada rede vêm dos rascunhos da própria V2 (E-mail 3 da
  mesma rede; no E-mail 4, link S2D02 e UTM EMAIL4, iguais ao rascunho V10 do E-mail 4).
- Supermercado da pessoa no texto principal dos E-mails 3 e 4 (NO_MERCADO), escrito por rede.
- O segundo bloco segue o padrão dos E-mails 1 e 2 da V2: o texto do campo de banner vira a abertura
  do bloco, os benefícios viram blocos do content_blocks.py e o texto complementar fecha o bloco.
  Formato varia por e-mail para a régua não repetir: E-mail 1 grade, E-mail 2 números, E-mail 3 lista,
  E-mail 4 grade com o Visa primeiro.
- O texto reorganizado fica escrito abaixo (EDITORIAL); check_facts confere que cada número e
  benefício citado existe na planilha, para não sair benefício novo.
- <v13_atual.html> é o template gravado hoje no slot builtin-plurix-v13 e <v13_novo.html> o
  PLURIX_V13_TEMPLATE do fixture (segundo bloco em largura total e faixa alternando amarelo/bege).
  A migration só troca os trechos que mudaram, confere o md5 antes e depois, e confere o md5 de cada
  campo gravado nos briefings: qualquer diferença desfaz a migration inteira.
"""
import difflib
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_topo_dinamico_rows import name_token, plain, runs, to_html  # noqa: E402
from content_blocks import benefit_grid, check_list, check_size, closing, lead  # noqa: E402

SEGMENT = "V2 TOPO DE FUNIL DINÂMICO"
RULER_ID = "d2d20000-0000-4000-8000-000000000001"
GROUP_E4 = "d2d20000-0000-4000-8000-0000000000e4"
TEMPLATE = "builtin-plurix-v13"
SEQUENCES = ["E-mail 3", "E-mail 4"]
NETWORKS = ["AMIGAO", "AVENIDA", "BOA", "COMPRE MAIS", "PARANA", "SUPERPAO"]

# Supermercado da pessoa no texto. Cada linha do briefing é de uma rede e o template escolhe a linha
# pela INVESTIDA da audiência, então o nome vai escrito por rede (com acento e caixa certos), em vez de
# %%INVESTIDA%%, que traz o código ("AMIGAO"). Compre Mais e Superpão não podem ser citados (transição
# de marca, mesma regra do assunto e da assinatura): ficam com o genérico.
NO_MERCADO = {
    "AMIGAO": "no Amigão", "AVENIDA": "no Avenida", "BOA": "no Boa", "PARANA": "no Paraná",
    "COMPRE MAIS": "no seu supermercado", "SUPERPAO": "no seu supermercado",
}
APPROVAL = "aprovado pelo Pablo em 15/10/2026"

# Saudação da faixa de limite: uma por e-mail.
LIMIT_MESSAGE = {
    "E-mail 2": "Seu cartão está mais perto, {{nome}}: você tem R$ {{limite}} de limite pré-aprovado no cartão +amigo.",
    "E-mail 3": "Ainda dá tempo, {{nome}}: você tem R$ {{limite}} de limite pré-aprovado para pedir seu cartão +amigo.",
    "E-mail 4": "Vale lembrar, {{nome}}: você tem R$ {{limite}} de limite pré-aprovado no cartão +amigo.",
}

# E-mail 1: mesmo bloco já aprovado (reproduzido byte a byte pelo content_blocks), agora com o
# número de cada quadro em negrito.
EMAIL1_COPY2 = (
    lead("E para aproveitar benefícios em diferentes momentos da rotina, conheça algumas das vantagens de ter seu <b>cartão +amigo</b>:")
    + benefit_grid([
        ("Descontos extras no mercado", "<b>A partir de 5%</b> em produtos selecionados, pagando com o cartão +amigo."),
        ("Mais prazo para pagar", "<b>Até 40 dias</b>, conforme a data da compra e o vencimento da fatura."),
        ("Benefícios com Visa", "Acesso a promoções, descontos, sorteios e <b>benefícios exclusivos</b>."),
        ("Vantagens para sua saúde", "<b>Até 70% de desconto</b> em consultas, exames e medicamentos pelo Programa Você Bem Saúde."),
    ], allow_html=True)
    + closing("Peça seu cartão +amigo e descubra essas vantagens!")
)

# Segundo bloco reorganizado, sem mudar o sentido da planilha.
EDITORIAL = {
    "E-mail 3": {
        "TITULO_COPY_2": "Mais motivos para aproveitar as compras",
        # Ajuste fino: o desconto em produtos selecionados já está no texto principal e na lista.
        "copy1": [("em produtos selecionados do supermercado.", "em produtos selecionados {no_mercado}.")],
        "lead": "Com o <b>cartão de crédito do Clube +amigo</b>, você paga suas compras com segurança e praticidade "
                "e ainda conta com <b>benefícios no dia a dia</b>:",
        "block": lambda: check_list([
            "<b>A partir de 5% de desconto</b> em produtos selecionados.",
            "<b>Até 40 dias para pagar</b>: mais flexibilidade para você.",
            "<b>Até 70% de desconto</b> em consultas, exames e medicamentos.",
            "<b>Benefícios com Visa</b>: aceito no Brasil e no mundo, com acesso a promoções, descontos, sorteios e benefícios exclusivos.",
        ]),
    },
    "E-mail 4": {
        "TITULO_COPY_2": "Vantagens para as compras e para o dia a dia",
        # Ajuste fino: promoções e sorteios do Visa ficam só no cartão da grade, sem repetir no texto.
        "copy1": [
            ("do que as compras no supermercado.", "do que as compras {no_mercado}."),
            ("ele é aceito em estabelecimentos no Brasil e no exterior, além de dar acesso a <b>promoções, descontos, "
             "sorteios e benefícios exclusivos</b>. E você ainda pode aproveitar <b>benefícios extras para economizar</b> "
             "e cuidar da sua rotina.",
             "ele é aceito em estabelecimentos no Brasil e no exterior e ainda traz <b>benefícios extras para você "
             "economizar</b> e cuidar da sua rotina."),
        ],
        "lead": "Com o <b>cartão de crédito +amigo</b>, você conta com:",
        "block": lambda: benefit_grid([
            ("Visa no Brasil e no mundo", "Mais possibilidades de uso e o <b>Programa Vai de Visa</b>, com promoções, descontos, sorteios e benefícios exclusivos."),
            ("Desconto extra no mercado", "<b>+5% de desconto extra</b> em produtos selecionados nas redes participantes."),
            ("Mais prazo para pagar", "<b>Até 40 dias</b>, conforme a data da compra e o vencimento da fatura."),
            ("Desconto em saúde", "<b>Até 70% de desconto</b> em consultas, exames e medicamentos pelo Programa Você Bem."),
        ], allow_html=True),
    },
}

# Fatos que cada e-mail cita: precisam existir no texto da planilha daquele e-mail.
FACTS = {
    "E-mail 3": ["5%", "70%", "40 dias", "Visa", "Brasil e no mundo", "promoções, descontos, sorteios e benefícios exclusivos",
                 "segurança e praticidade", "benefícios no dia a dia", "consultas, exames e medicamentos"],
    "E-mail 4": ["Visa", "Brasil e no mundo", "Vai de Visa", "promoções, descontos, sorteios e benefícios exclusivos",
                 "5%", "redes participantes", "40 dias", "vencimento da fatura", "70%", "Você Bem",
                 "consultas, exames e medicamentos"],
}


def check_facts(sequence, source_text):
    flat = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", source_text))
    missing = [fact for fact in FACTS[sequence] if fact.lower() not in flat.lower()]
    if missing:
        raise ValueError(f"{sequence}: fatos sem origem na planilha: {missing}")


def read_sheet(xlsx):
    ws = openpyxl.load_workbook(xlsx, rich_text=True)["Briefing Fábrica"]
    columns = {str(ws.cell(4, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    fields = {}
    for r in range(9, 47):
        field = (ws.cell(r, 3).value or "").strip()
        if field:
            fields[field] = r
    return ws, columns, fields


def build(xlsx):
    ws, columns, fields = read_sheet(xlsx)
    week_row = next(r for r in range(1, 9) if str(ws.cell(r, 2).value or "").strip() == "Semana")
    out = []
    for sequence in SEQUENCES:
        col = columns[sequence]
        get = lambda field: runs(ws.cell(fields[field], col))  # noqa: E731
        assert plain(get("SEQUENCIA")) == sequence, f"coluna {sequence} fora do lugar"
        banner, title2, copy2 = (to_html(get(f)) for f in ("BANNER_1_CORPO", "TITULO_COPY_2", "COPY_2_PRETO"))
        check_facts(sequence, " ".join([to_html(get("COPY_1_PRETO")), banner, title2, copy2]))
        ed = EDITORIAL[sequence]
        copy2_html = lead(ed["lead"]) + ed["block"]() + closing(re.sub(r"</?b>", "", copy2))
        check_size("COPY_2_PRETO", copy2_html)
        copy1 = name_token(to_html(get("COPY_1_PRETO")))
        for old, new in ed["copy1"]:
            assert copy1.count(old) == 1, f"{sequence}: trecho do texto principal mudou na planilha: {old!r}"
            copy1 = copy1.replace(old, new)
        out.append({
            "sequence": sequence,
            "week": plain(runs(ws.cell(week_row, col))),
            "by_network": {net: {"COPY_1_PRETO": copy1.replace("{no_mercado}", place)} for net, place in NO_MERCADO.items()},
            "data": {
                "TP_CAMPANHA": plain(get("TP_CAMPANHA")),
                "SEQUENCIA": sequence,
                "ASSUNTO": plain(get("ASSUNTO")),
                "PRE_CABECALHO": plain(get("PRE_CABECALHO")),
                "HEADER_TITULO": plain(get("TB_HEADER_VARIACOES · TITULO")),
                "HEADER_SUBTITULO": plain(get("TB_HEADER_VARIACOES · SUBTITULO")),
                "HEADER_IMAGEM_CARTAO": plain(get("TB_HEADER_VARIACOES · IMAGEM_CARTAO")),
                "MENSAGEM_LIMITE": LIMIT_MESSAGE[sequence],
                "TITULO_COPY_1_AZUL": name_token(plain(get("TITULO_COPY_1_AZUL"))),
                "TITULO_CTA_1": plain(get("TITULO_CTA_1")).upper(),
                "BANNER_1_CORPO": "",
                "TITULO_COPY_2": ed["TITULO_COPY_2"],
                "COPY_2_PRETO": copy2_html,
                "TITULO_CTA_2": plain(get("TITULO_CTA_2")).upper(),
                "BANNER_2_CORPO": "",
                "TAMANHO_DA_FONTE_TITULO_COPY_2": "24",
            },
        })
    return out


def tag(text):
    assert "$gaas$" not in text
    return f"$gaas${text}$gaas$"


def jsonb(value):
    return f"{tag(json.dumps(value, ensure_ascii=False))}::jsonb"


def by_network_values(by_network):
    return ",\n    ".join(f"({tag(net)}, {jsonb(patch)})" for net, patch in by_network.items())


def update(sequence, patch, summary, by_network=None):
    """Atualiza as 6 redes de um e-mail da V2 (patch comum + patch da rede), sobe a versão e grava o histórico."""
    by_network = by_network or {net: {} for net in NETWORKS}
    return f"""-- {sequence}: {summary}
with rede(nm, patch) as (values
    {by_network_values(by_network)}
), alvo as (
  select b.id, b.briefing_data || {jsonb(patch)} || rede.patch as dados
  from public.dynamic_email_briefings b
  join rede on rede.nm = b.briefing_data->>'NM_PRODUTO_INTERNO'
  where b.segment = {tag(SEGMENT)} and b.status <> 'archived' and b.briefing_data->>'SEQUENCIA' = {tag(sequence)}
), atualizados as (
  update public.dynamic_email_briefings b set briefing_data = alvo.dados, version = b.version + 1, updated_at = now()
  from alvo where b.id = alvo.id and b.briefing_data is distinct from alvo.dados
  returning b.id, b.version, b.briefing_data
)
insert into public.dynamic_email_briefing_versions (briefing_id, version, snapshot, change_summary, warnings)
select id, version, briefing_data, {tag(summary)}, '[]'::jsonb from atualizados
on conflict (briefing_id, version) do nothing;
"""


def md5(text):
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def template_patch(old, new):
    """Trechos do template que mudaram, como pares (antes, depois), cada antes único no template atual."""
    a, b = old.splitlines(keepends=True), new.splitlines(keepends=True)
    pairs = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == "equal":
            continue
        # Ancora na linha anterior para inserções puras e para garantir que o trecho seja único.
        i0, j0 = (i1 - 1, j1 - 1) if i1 > 0 else (i1, j1)
        before, after = "".join(a[i0:i2]), "".join(b[j0:j2])
        assert old.count(before) == 1, f"trecho do template não é único: {before[:80]!r}"
        pairs.append((before, after))
    patched = old
    for before, after in pairs:
        patched = patched.replace(before, after)
    assert patched == new, "as trocas não reproduzem o template novo"
    return pairs


def template_sql(old, new):
    expr = "source"
    for before, after in template_patch(old, new):
        expr = f"replace({expr}, {tag(before)}, {tag(after)})"
    return f"""update public.dynamic_email_template_slots set source = {expr}, version = version + 1, updated_at = now()
where id = {tag(TEMPLATE)} and md5(source) = '{md5(old)}';"""


def checks(emails):
    """Confere o que foi gravado; qualquer diferença desfaz a migration inteira."""
    e3, e4 = emails
    expected = [("E-mail 1", "COPY_2_PRETO", None, EMAIL1_COPY2), ("E-mail 2", "MENSAGEM_LIMITE", None, LIMIT_MESSAGE["E-mail 2"])]
    for item in emails:
        expected.append((item["sequence"], "COPY_2_PRETO", None, item["data"]["COPY_2_PRETO"]))
        expected.append((item["sequence"], "MENSAGEM_LIMITE", None, item["data"]["MENSAGEM_LIMITE"]))
        for net, patch in item["by_network"].items():
            expected.append((item["sequence"], "COPY_1_PRETO", net, patch["COPY_1_PRETO"]))
    lines = []
    for sequence, field, net, value in expected:
        scope = f" and briefing_data->>'NM_PRODUTO_INTERNO' = {tag(net)}" if net else ""
        label = f"{sequence} {field}{' ' + net if net else ''}"
        lines.append(f"""  if exists (select 1 from public.dynamic_email_briefings where segment = {tag(SEGMENT)} and status <> 'archived'
      and briefing_data->>'SEQUENCIA' = {tag(sequence)}{scope} and md5(coalesce(briefing_data->>{tag(field)}, '')) <> '{md5(value)}')
    or not exists (select 1 from public.dynamic_email_briefings where segment = {tag(SEGMENT)} and status <> 'archived'
      and briefing_data->>'SEQUENCIA' = {tag(sequence)}{scope}) then
    raise exception 'V2 E-mails 3 e 4: {label} diferente do gerado';
  end if;""")
    body = "\n".join(lines)
    return f"""do $check$
begin
  if (select md5(source) from public.dynamic_email_template_slots where id = {tag(TEMPLATE)}) <> '{TEMPLATE_MD5_PLACEHOLDER}' then
    raise exception 'PLURIX V13: template gravado diferente do fixture';
  end if;
  if (select count(*) from public.dynamic_email_briefings where segment = {tag(SEGMENT)} and status <> 'archived'
      and briefing_data->>'SEQUENCIA' = 'E-mail 4' and briefing_data->>'LINK_CTA_1' like '%S2D02%'
      and briefing_data->>'UTM_CAMPANHA' like '%\\_EMAIL4\\_%') <> 6 then
    raise exception 'V2 E-mail 4: esperado 6 redes com link S2D02 e UTM EMAIL4';
  end if;
{body}
end
$check$;"""


TEMPLATE_MD5_PLACEHOLDER = "{template_md5}"


def sql(emails, old_template, new_template):
    e3, e4 = emails
    e4_strategy = jsonb({"origem": "planilha Regua_Plurix_TopoFunil_Proposta_Outubro26_3, aba Briefing Fábrica",
                         "rede": "links, UTM, assinatura e nota legal do E-mail 3 da V2"})
    e4_source = {"origem": "planilha Regua_Plurix_TopoFunil_Proposta_Outubro26_3 (E-mail 4)",
                 "mudanca": "segundo bloco em blocos visuais (content_blocks.py)",
                 "rede": "E-mail 3 da V2 da mesma rede, link S2D02 e UTM EMAIL4"}
    return f"""-- V2 TOPO DE FUNIL DINÂMICO: E-mails 3 e 4 da planilha de outubro, mais os ajustes de régua aprovados junto.
-- 1. PLURIX V13: segundo bloco sem imagem em largura total (E-mail 3 em diante) e faixa de limite
--    alternando amarelo (#E6A813, e-mails ímpares) e bege (#EBE0CF, pares). Mesmo texto do fixture.
-- 2. E-mail 1: número de cada quadro em negrito.  3. E-mail 2: volta a saudação própria da faixa.
-- 4. E-mail 3: segundo bloco preenchido (lista com check).  5. E-mail 4: novo, 6 redes, a partir do E-mail 3.
-- Gerado por scripts/email-factory/build_topo_dinamico_v2_e3e4.py.

{template_sql(old_template, new_template)}

{update("E-mail 1", {"COPY_2_PRETO": EMAIL1_COPY2}, "Número de cada quadro em negrito (E-mail 1).")}
{update("E-mail 2", {"MENSAGEM_LIMITE": LIMIT_MESSAGE["E-mail 2"]}, "Saudação da faixa volta a ser própria do E-mail 2.")}
{update("E-mail 3", e3["data"], "Conteúdo da planilha de outubro; segundo bloco em lista com check; supermercado no texto.", e3["by_network"])}
-- E-mail 4: uma linha por rede, copiada do E-mail 3 da mesma rede (ids ...0000012N -> ...0000013N).
insert into public.dynamic_email_email_strategies (ruler_strategy_id, campaign_group_id, partner, segment, week_key, sequence, subject, preheader, technical_status, editorial_status, visual_status, certification_status, field_provenance, updated_by_type, update_source, change_reason, llm_model, functional_name)
values ({tag(RULER_ID)}::uuid, {tag(GROUP_E4)}::uuid, 'Plurix', {tag(SEGMENT)}, {tag(e4["week"])}, 'E-mail 4',
  {tag(e4["data"]["ASSUNTO"])}, {tag(e4["data"]["PRE_CABECALHO"])}, 'draft', 'draft', 'needs_review', 'not_tested',
  {e4_strategy}, 'llm', 'llm', {tag("E-mail 4 da planilha de outubro em blocos visuais, " + APPROVAL)}, 'claude-opus-5-5',
  'V2 TOPO DINÂMICO · E-mail 4')
on conflict (campaign_group_id) do nothing;

with rede(nm, patch) as (values
    {by_network_values(e4["by_network"])}
), origem as (
  select b.*, rede.patch, ('d2d20000-0000-4000-8000-0000000001' || '3' || right(b.id::text, 1))::uuid as novo_id
  from public.dynamic_email_briefings b
  join rede on rede.nm = b.briefing_data->>'NM_PRODUTO_INTERNO'
  where b.segment = {tag(SEGMENT)} and b.status <> 'archived' and b.briefing_data->>'SEQUENCIA' = 'E-mail 3'
)
insert into public.dynamic_email_briefings (id, briefing_data, partner, segment, subgroup, week_key, activity_names, campaign_group_id, status, version, template_slot_id, journey_confirmed, import_source)
select o.novo_id,
  o.briefing_data || {jsonb(e4["data"])} || o.patch || jsonb_build_object(
    '__id', o.novo_id::text,
    'LINK_CTA_1', replace(o.briefing_data->>'LINK_CTA_1', 'S2D01', 'S2D02'),
    'LINK_CTA_2', replace(o.briefing_data->>'LINK_CTA_2', 'S2D01', 'S2D02'),
    'LINK_BANNER_1_CORPO', replace(o.briefing_data->>'LINK_BANNER_1_CORPO', 'S2D01', 'S2D02'),
    'LINK_BANNER_2_CORPO', replace(o.briefing_data->>'LINK_BANNER_2_CORPO', 'S2D01', 'S2D02'),
    'UTM_CAMPANHA', replace(o.briefing_data->>'UTM_CAMPANHA', '_EMAIL3_', '_EMAIL4_')),
  o.partner, o.segment, o.subgroup, {tag(e4["week"])}, array[]::text[], {tag(GROUP_E4)}::uuid, 'draft', 1, {tag(TEMPLATE)}, false,
  {jsonb(e4_source)} || jsonb_build_object('importKey', {tag(SEGMENT + "|E-mail 4|")} || (o.briefing_data->>'NM_PRODUTO_INTERNO'))
from origem o
on conflict (id) do nothing;

insert into public.dynamic_email_briefing_versions (briefing_id, version, snapshot, change_summary, warnings)
select b.id, 1, b.briefing_data, 'E-mail 4 criado a partir da planilha de outubro e do E-mail 3 da mesma rede.', '[]'::jsonb
from public.dynamic_email_briefings b
where b.segment = {tag(SEGMENT)} and b.campaign_group_id = {tag(GROUP_E4)}::uuid
on conflict (briefing_id, version) do nothing;

{checks(emails).replace(TEMPLATE_MD5_PLACEHOLDER, md5(new_template))}
"""


def main(xlsx, old_template_path, new_template_path, out_json, out_sql=None):
    emails = build(xlsx)
    json.dump({"emails": emails, "email1_copy2": EMAIL1_COPY2, "limit_messages": LIMIT_MESSAGE},
              open(out_json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    if out_sql:
        read = lambda path: open(path, encoding="utf-8").read()  # noqa: E731
        open(out_sql, "w", encoding="utf-8").write(sql(emails, read(old_template_path), read(new_template_path)))
        print(out_sql)
    for item in emails:
        print(f"\n=== {item['sequence']} ({item['week']})")
        for key, value in item["data"].items():
            print(f"{key}: {value if key != 'COPY_2_PRETO' else f'<{len(value)} caracteres>'}")


if __name__ == "__main__":
    main(*sys.argv[1:6])
