# -*- coding: utf-8 -*-
"""TOPO DE FUNIL DINÂMICO: lê a aba "Briefing Fábrica" (linhas 9 a 46, E-mail 1 a 3) e monta
as 18 linhas de briefing (3 e-mails x 6 redes) no formato da DE.

Uso: python scripts/email-factory/build_topo_dinamico_rows.py <planilha.xlsx> <v10_e1e3.json> <saida.json>

- Lê por rótulo (coluna "Onde fica"), não por posição de linha.
- Negrito do Excel vira <b>, quebra de linha vira <br>, lista "- " vira "• ".
- Linhas marcadas como automáticas não entram (versões por rede, logo, pool, assinatura).
- Campos de cada rede (links, UTM, assinatura, rodapé, nota legal, cores e tamanhos) vêm dos
  rascunhos E1 a E3 da régua V10, que ocupam as mesmas posições da régua.
"""
import json
import re
import sys

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

EMAILS = {"E-mail 1": 5, "E-mail 2": 6, "E-mail 3": 7}
NBSP = "\xa0"
NETWORKS = ["AMIGAO", "AVENIDA", "BOA", "COMPRE MAIS", "PARANA", "SUPERPAO"]
# Ajuste editorial aprovado: erro de digitação da planilha.
TYPO_FIXES = {"aproveitar as melhoras ": "aproveitar as melhores "}
# Campos que continuam vindo do rascunho V10 de cada rede.
FROM_V10 = [
    "UTM_CAMPANHA", "LINK_CTA_1", "LINK_CTA_2", "LINK_BANNER_1_CORPO", "LINK_BANNER_2_CORPO",
    "BANNER_3_CORPO", "LINK_BANNER_3_CORPO", "RODAPE", "NOTA_LEGAL", "COR_NOTA_LEGAL",
    "TAMANHO_DA_FONTE_NOTA_LEGAL", "CARTAO_NM_COMERCIAL", "COR_COPY_1", "TAMANHO_DA_FONTE_TITULO_COPY_1",
    "COR_COPY_PRETO_1", "TAMANHO_DA_FONTE_TITULO_COPY_PRETO_1", "COR_TITULO_COPY_2",
    "TAMANHO_DA_FONTE_TITULO_COPY_2", "COR_COPY_2", "TAMANHO_DA_FONTE_COPY_2", "DT_FIM",
]


def runs(cell):
    value = cell.value
    if value is None:
        return []
    whole = bool(cell.font and cell.font.b)
    if isinstance(value, CellRichText):
        out = []
        for part in value:
            if isinstance(part, TextBlock):
                out.append((part.text, bool(part.font and part.font.b)))
            else:
                out.append((str(part), whole))
        return out
    return [(str(value), whole)]


def to_html(parts):
    """Negrito -> <b>, quebra -> <br>, "- " no começo da linha -> "• "."""
    pieces = []
    for text, bold in parts:
        text = text.replace(NBSP, " ").replace("\r", "")
        for index, chunk in enumerate(text.split("\n")):
            if index:
                pieces.append("<br>")
            if not chunk:
                continue
            if bold and chunk.strip():
                lead = chunk[: len(chunk) - len(chunk.lstrip())]
                tail = chunk[len(chunk.rstrip()):]
                pieces.append(f"{lead}<b>{chunk.strip()}</b>{tail}")
            else:
                pieces.append(chunk)
    html = "".join(pieces)
    html = re.sub(r"</b>(\s*)<b>", r"\1", html)          # junta negritos vizinhos
    html = re.sub(r"(^|<br>)(\s*)(<b>)?-\s+", lambda m: f"{m.group(1)}{m.group(3) or ''}• ", html)
    html = re.sub(r"[ \t]{2,}", " ", html).strip()
    html = re.sub(r"(<br>\s*)+$", "", html)
    for wrong, right in TYPO_FIXES.items():
        html = html.replace(wrong, right)
    return html


def plain(parts):
    return re.sub(r"\s+", " ", "".join(t for t, _ in parts).replace(NBSP, " ")).strip()


def name_token(text):
    return re.sub(r"%{0,2}PRI_NOME%%", "%%=v(@FirstName)=%%", text)


def main(xlsx, v10_json, out):
    ws = openpyxl.load_workbook(xlsx, rich_text=True)["Briefing Fábrica"]
    by_field = {}
    for r in range(9, 47):
        field = (ws.cell(r, 3).value or "").strip()
        if field:
            by_field[field] = r
    get = lambda field, col: runs(ws.cell(by_field[field], col))

    v10 = json.load(open(v10_json, encoding="utf-8"))
    v10_by_key = {(x["data"]["SEQUENCIA"], x["data"]["NM_PRODUTO_INTERNO"]): x for x in v10}

    limit_text = plain(get("LIMITE (audiência)", 5))
    match = re.search(r"“(.+?)”", limit_text)
    message = match.group(1).replace("[nome]", "{{nome}}").replace("[limite]", "{{limite}}") if match else ""

    rows, report = [], []
    for sequence, col in EMAILS.items():
        week = plain(runs(ws.cell(5, col)))
        banner_text = to_html(get("BANNER_1_CORPO", col))
        copy2 = to_html(get("COPY_2_PRETO", col))
        # Texto escrito no campo de banner vira o começo do segundo bloco (decisão de 15/10/2026).
        copy2_full = "<br><br>".join(x for x in [banner_text, copy2] if x)
        shared = {
            "TP_CAMPANHA": plain(get("TP_CAMPANHA", col)),
            "SEQUENCIA": sequence,
            "ASSUNTO": plain(get("ASSUNTO", col)),
            "PRE_CABECALHO": plain(get("PRE_CABECALHO", col)),
            "HEADER": "",
            "HEADER_TITULO": plain(get("TB_HEADER_VARIACOES · TITULO", col)),
            "HEADER_SUBTITULO": plain(get("TB_HEADER_VARIACOES · SUBTITULO", col)),
            "HEADER_IMAGEM_CARTAO": plain(get("TB_HEADER_VARIACOES · IMAGEM_CARTAO", col)),
            "HEADER_LOGO": plain(get("HEADER (briefing)", col)),
            "MENSAGEM_LIMITE": message,
            "TITULO_COPY_1_AZUL": name_token(plain(get("TITULO_COPY_1_AZUL", col))),
            "COPY_1_PRETO": name_token(to_html(get("COPY_1_PRETO", col))),
            "TITULO_CTA_1": plain(get("TITULO_CTA_1", col)),
            "BANNER_1_CORPO": "",
            "TITULO_COPY_2": to_html(get("TITULO_COPY_2", col)),
            "COPY_2_PRETO": copy2_full,
            "TITULO_CTA_2": plain(get("TITULO_CTA_2", col)),
            "BANNER_2_CORPO": plain(get("BANNER_2_CORPO", col)),
            "DT_INICIO": "2026-10-15T00:00",
        }
        report.append({"sequence": sequence, "week": week, **shared})
        for network in NETWORKS:
            source = v10_by_key[(sequence, network)]
            row = {**{k: source["data"].get(k, "") for k in FROM_V10}, **shared, "NM_PRODUTO_INTERNO": network}
            rows.append({"sequence": sequence, "week": week, "network": network, "subgroup": source["subgroup"],
                         "v10_id": source["id"], "data": row})

    json.dump({"rows": rows, "report": report}, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    for item in report:
        print(f"\n=== {item['sequence']} ({item['week']})")
        for key in ["ASSUNTO", "PRE_CABECALHO", "HEADER_TITULO", "HEADER_SUBTITULO", "HEADER_IMAGEM_CARTAO", "HEADER_LOGO",
                    "MENSAGEM_LIMITE", "TITULO_COPY_1_AZUL", "COPY_1_PRETO", "TITULO_CTA_1", "TITULO_COPY_2", "COPY_2_PRETO", "TITULO_CTA_2"]:
            print(f"{key}: {item[key]}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
