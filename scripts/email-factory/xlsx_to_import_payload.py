# -*- coding: utf-8 -*-
"""Lê a aba editorial de uma régua (ex.: "TOPO CRM 3") e gera o payload fabrica-xlsx-import/1.

Uso:
  python scripts/email-factory/xlsx_to_import_payload.py <arquivo.xlsx> "<aba>" <saida.json> [--reader nome-do-modelo]

A leitura é por rótulo, nunca por posição fixa:
- cabeçalhos "E-mail N" definem os toques (extensível além de dois);
- a coluna "Campo (DE)" diz qual coluna do briefing a linha preenche;
- a coluna "Tipo" diz a origem: "Marketing escreve" = xlsx, "Cadastro governado" e
  "Identidade por rede" = governed (com a referência da coluna "Origem");
- seção com cabeçalho "Rede" = valores por rede; "(vazio)" = vazio de propósito,
  célula em branco = usa o valor comum;
- seção com cabeçalho "Severidade" = pendências.

O resultado é validado no GaaS por src/modules/dynamic-email/domain/xlsxImport.ts.
"""
import hashlib
import json
import re
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

import openpyxl

CONTRACT = "fabrica-xlsx-import/1"
EMPTY_MARK = "(vazio)"
TOUCH = re.compile(r"^E-mail \d+(\.\d+)?$")
RULER_FIELDS = {
    "ruler.name": "name", "ruler.segment": "segment", "TP_CAMPANHA": "tpCampanha",
    "ruler.businessClassification": "businessClassification", "ruler.partner": "partner",
    "ruler.product": "product", "ruler.templateSlotId": "templateSlotId",
}


def text(value):
    return "" if value is None else str(value).strip()


def read(path: Path, sheet_name: str, reader: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    headers = None
    ruler, touches, pendencies = {}, {}, []
    order = []

    def cell_ref(row, col):
        return f"{sheet_name}!{openpyxl.utils.get_column_letter(col)}{row}"

    for row in range(1, ws.max_row + 1):
        values = [text(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        if not any(values):
            continue
        # Linha de cabeçalho: guarda a posição de cada rótulo.
        if "Campo (DE)" in values or "Campo" in values:
            headers = {label: index + 1 for index, label in enumerate(values) if label}
            for label in values:
                if TOUCH.match(label) and label not in order:
                    order.append(label)
                    touches[label] = {"sequence": label, "weekKey": "", "shared": {}, "variants": {}}
            continue
        if not headers:
            continue
        get = lambda label: text(ws.cell(row, headers[label]).value) if label in headers else ""
        first = next(iter(headers))
        if first == "Item":  # seção da régua
            field = get("Campo")
            if field in RULER_FIELDS:
                ruler[RULER_FIELDS[field]] = get("Valor")
            continue
        if first == "Severidade":
            message = get("Pendência")
            if message:
                column = get("Campo (DE)")
                pendencies.append({
                    "severity": "blocker" if get("Severidade").lower().startswith("bloq") else "review",
                    "scope": get("Escopo") or "régua",
                    **({"field": column} if re.fullmatch(r"[A-Z0-9_]+", column) else {}),
                    "message": message,
                })
            continue
        column = get("Campo (DE)")
        kind = get("Tipo")
        if first == "Rede":
            network_rows = True
        else:
            network_rows = False
        for touch in order:
            col = headers.get(touch)
            if not col:
                continue
            raw = text(ws.cell(row, col).value)
            if column == "(não entra na DE)":
                if get("Nome na tela") == "Semana":
                    touches[touch]["weekKey"] = raw
                continue
            if not re.fullmatch(r"[A-Z0-9_]+", column):
                continue
            if column in ("SEQUENCIA", "NM_PRODUTO_INTERNO", "TP_CAMPANHA"):
                if network_rows and column == "NM_PRODUTO_INTERNO":
                    touches[touch]["variants"].setdefault(get("Rede"), {"network": raw, "fields": {}})["network"] = raw
                continue  # identidade técnica é derivada da régua e do toque
            if network_rows and not raw:
                continue  # branco = usa o valor comum
            value = "" if raw == EMPTY_MARK else raw
            if kind == "Marketing escreve":
                field = {"value": value, "origin": "xlsx", "source": cell_ref(row, col)}
            else:
                field = {"value": value, "origin": "governed", "source": f"{get('Origem') or kind} ({cell_ref(row, col)})"}
            if network_rows:
                touches[touch]["variants"].setdefault(get("Rede"), {"network": "", "fields": {}})["fields"][column] = field
            else:
                touches[touch]["shared"][column] = field

    payload = {
        "contract": CONTRACT,
        "status": "draft",
        "source": {
            "file": path.name,
            "sheet": sheet_name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "readAt": datetime.now(timezone(timedelta(hours=-3))).isoformat(timespec="seconds"),
            "reader": reader,
        },
        "ruler": ruler,
        "touches": [
            {**{k: v for k, v in touches[t].items() if k != "variants"}, "variants": list(touches[t]["variants"].values())}
            for t in order
        ],
        "pendencies": pendencies,
    }
    return payload


if __name__ == "__main__":
    args = sys.argv[1:]
    reader = "manual"
    if "--reader" in args:
        index = args.index("--reader")
        reader = args[index + 1]
        del args[index:index + 2]
    if len(args) != 3:
        sys.exit(__doc__)
    result = read(Path(args[0]), args[1], reader)
    Path(args[2]).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{args[2]}: {len(result['touches'])} toques, "
          f"{sum(len(t['variants']) for t in result['touches'])} variantes, {len(result['pendencies'])} pendências")
