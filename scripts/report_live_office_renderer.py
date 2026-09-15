#!/usr/bin/env python3
"""Renderiza uma prova vertical do Report Live em PPTX e PDF sem Google.

O renderer consome somente o artefato imutável certificado. Regras de negócio,
narrativa, seleção de dados e geometria continuam pertencendo ao engine TS.
"""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any, Iterable


CANVAS_WIDTH_PT = 720.0
CANVAS_HEIGHT_PT = 405.0
AFINZ = {
    "black": "000000",
    "white": "FFFFFF",
    "cyan": "00C6CC",
    "lime": "D3FF00",
    "red": "E74742",
}
DEFAULT_SLIDES = ("c0", "c4", "p1_serasa")


def validate_artifact_identity(
    artifact: dict[str, Any], expected_run_id: str, expected_content_hash: str,
) -> None:
    if artifact.get("run_id") != expected_run_id:
        raise ValueError("run_id do artefato diverge da resposta assinada")
    if artifact.get("content_hash") != expected_content_hash:
        raise ValueError("content_hash do artefato diverge da resposta assinada")


def records(table: list[list[Any]] | None) -> list[dict[str, Any]]:
    if not table:
        return []
    headers = [str(value) for value in table[0]]
    return [dict(zip(headers, row)) for row in table[1:]]


def clean_text(value: Any) -> str:
    return str(value or "").replace("\u00a0", " ").replace("\ufffd", "-")


def parse_number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def compact_number(value: float | int | None) -> str:
    if value is None:
        return "indisponível"
    absolute = abs(float(value))
    if absolute >= 1_000_000:
        return f"{value / 1_000_000:.1f} mi".replace(".", ",")
    if absolute >= 1_000:
        return f"{value / 1_000:.1f} mil".replace(".", ",")
    return f"{value:,.0f}".replace(",", ".")


def month_label(period_start: str) -> str:
    months = [
        "janeiro", "fevereiro", "março", "abril", "maio", "junho",
        "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
    ]
    year, month, _ = period_start.split("-")
    return f"{months[int(month) - 1]} de {year}"


def normalize_narrative(value: str) -> str:
    text = clean_text(value)
    text = re.sub(r"^LEITURA DA DECISÃO\s*", "", text, flags=re.IGNORECASE)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def build_render_plan(
    artifact: dict[str, Any],
    slide_ids: Iterable[str] = DEFAULT_SLIDES,
) -> dict[str, Any]:
    if artifact.get("artifact_version") != "1":
        raise ValueError("Versão de artefato não suportada")
    if artifact.get("versions", {}).get("spec") != "3.0":
        raise ValueError("A prova vertical exige o contrato editorial 3.0")

    tabs = artifact.get("tabs", {})
    registry = {row["slide_instance_id"]: row for row in records(tabs.get("VIEW_REGISTRY"))}
    slides = {row["slide_instance_id"]: row for row in artifact.get("slides", [])}
    blueprints = {
        row["slide_instance_id"]: row for row in artifact.get("slide_blueprints", [])
    }
    layouts = {
        row["slide_instance_id"]: row for row in records(tabs.get("VIEW_EDITORIAL_LAYOUTS"))
    }
    rulers = records(tabs.get("VIEW_EDITORIAL_RULERS"))

    planned: list[dict[str, Any]] = []
    for slide_id in slide_ids:
        if slide_id not in slides or slide_id not in blueprints or slide_id not in registry:
            raise ValueError(f"Slide {slide_id} ausente do artefato certificado")
        slide = slides[slide_id]
        blueprint = blueprints[slide_id]["blueprint"]
        item: dict[str, Any] = {
            "slide_instance_id": slide_id,
            "slide_code": slide["slide_code"],
            "title": clean_text(registry[slide_id]["title"]),
            "partner": clean_text(slide.get("partner")),
            "confidence": clean_text(slide.get("confidence_label")),
            "confidence_status": clean_text(slide.get("confidence_status")),
            "narrative": normalize_narrative(blueprint.get("narrative", "")),
            "archetype": blueprint["visual"]["archetype"],
            "geometry": blueprint["visual"]["geometry"],
            "source_view": slide.get("source_view"),
            "support_text": clean_text(layouts.get(slide_id, {}).get("support_text")),
        }
        if slide_id == "c0":
            item["manifest"] = {
                row["campo"]: row["valor"] for row in records(tabs.get("VIEW_RUN_MANIFEST"))
            }
        elif slide_id == "c4":
            pacing = records(tabs.get("VIEW_PACING_ISODAYS"))
            item["chart"] = {
                "categories": [int(row["day_of_period"]) for row in pacing],
                "series": [
                    {
                        "name": "Agosto de 2026",
                        "color": AFINZ["cyan"],
                        "values": [parse_number(row.get("cumulative_cards")) for row in pacing],
                    },
                    {
                        "name": "Julho de 2026",
                        "color": AFINZ["black"],
                        "values": [parse_number(row.get("previous_equivalent_cumulative_cards")) for row in pacing],
                    },
                ],
                "x_title": "Dia do mês",
                "y_title": "Cartões acumulados",
            }
            if pacing:
                current = parse_number(pacing[-1].get("cumulative_cards"))
                previous = parse_number(pacing[-1].get("previous_equivalent_cumulative_cards"))
                item["summary"] = {
                    "current": current,
                    "previous": previous,
                    "delta": None if current is None or not previous else current / previous - 1,
                }
        elif slide_id.startswith("p1_"):
            metric_rows = sorted(
                (row for row in rulers if row.get("slide_instance_id") == slide_id),
                key=lambda row: int(row.get("metric_order") or 0),
            )[:2]
            item["metrics"] = [
                {
                    "label": clean_text(row.get("metric_label")),
                    "value": clean_text(row.get("value_text")),
                    "delta": clean_text(row.get("delta_text")),
                    "range": clean_text(row.get("range_text")),
                    "verdict": clean_text(row.get("verdict_text") or row.get("verdict")),
                }
                for row in metric_rows
            ]
        planned.append(item)

    return {
        "run_id": artifact["run_id"],
        "period_start": artifact["period_start"],
        "period_end": artifact["period_end"],
        "report_profile": artifact["report_profile"],
        "versions": artifact["versions"],
        "slides": planned,
    }


def _remove_all_slides(prs: Any) -> None:
    slide_ids = prs.slides._sldIdLst  # python-pptx ainda não expõe remoção pública.
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _remove_placeholders(slide: Any) -> None:
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            shape._element.getparent().remove(shape._element)


def render_pptx(plan: dict[str, Any], template: Path, logo: Path, output: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    prs = Presentation(str(template))
    _remove_all_slides(prs)
    content_layout = prs.slide_layouts[18]
    sx = prs.slide_width / CANVAS_WIDTH_PT
    sy = prs.slide_height / CANVAS_HEIGHT_PT

    def rgb(hex_value: str) -> Any:
        return RGBColor.from_string(hex_value.lstrip("#"))

    def box(slide: Any, x: float, y: float, w: float, h: float, fill: str, line: str | None = None) -> Any:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, int(x * sx), int(y * sy), int(w * sx), int(h * sy)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(line or fill)
        return shape

    def text_box(
        slide: Any, text: str, x: float, y: float, w: float, h: float,
        size: float, color: str, bold: bool = False,
        align: Any = PP_ALIGN.LEFT, valign: Any = MSO_ANCHOR.TOP,
    ) -> Any:
        shape = slide.shapes.add_textbox(int(x * sx), int(y * sy), int(w * sx), int(h * sy))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Pt(2)
        frame.margin_top = frame.margin_bottom = Pt(1)
        frame.vertical_anchor = valign
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = clean_text(text)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        return shape

    def base_slide(background: str) -> Any:
        slide = prs.slides.add_slide(content_layout)
        _remove_placeholders(slide)
        box(slide, 0, 0, CANVAS_WIDTH_PT, CANVAS_HEIGHT_PT, background)
        return slide

    def footer(slide: Any, page: int, dark: bool = False) -> None:
        color = AFINZ["white"] if dark else AFINZ["black"]
        text_box(slide, f"Report Live. Run {plan['run_id'][:8]}. Spec {plan['versions']['spec']}.",
                 28, 374, 500, 16, 8, color)
        text_box(slide, str(page), 670, 374, 22, 16, 8, color, align=PP_ALIGN.RIGHT)

    for page, item in enumerate(plan["slides"], start=1):
        if item["slide_instance_id"] == "c0":
            slide = base_slide(AFINZ["black"])
            box(slide, 28, 66, 76, 5, AFINZ["cyan"])
            box(slide, 108, 66, 40, 5, AFINZ["lime"])
            text_box(slide, "report live", 28, 96, 500, 62, 34, AFINZ["white"], True)
            text_box(slide, month_label(plan["period_start"]), 28, 158, 500, 34, 20, AFINZ["cyan"], True)
            text_box(slide, "prova vertical em PowerPoint e PDF", 28, 204, 500, 28, 14, AFINZ["white"])
            manifest = item["manifest"]
            contract = (
                f"Perfil {manifest.get('report_profile', '')}. "
                f"Qualidade {manifest.get('quality_status', '')}. "
                f"Cutoff CRM {json.loads(manifest.get('source_cutoffs', '{}')).get('crm', 'indisponível')}."
            )
            text_box(slide, contract, 28, 262, 560, 50, 12, AFINZ["white"])
            box(slide, 600, 42, 80, 34, AFINZ["white"])
            slide.shapes.add_picture(str(logo), int(606 * sx), int(48 * sy), width=int(68 * sx))
            footer(slide, page, dark=True)
            continue

        slide = base_slide(AFINZ["white"])
        box(slide, 0, 0, CANVAS_WIDTH_PT, 8, AFINZ["black"])
        text_box(slide, item["title"].lower(), 28, 28, 570, 38, 24, AFINZ["black"], True)
        text_box(slide, f"Confiança {item['confidence']}. Fonte {item['source_view']}.",
                 28, 68, 570, 20, 9, AFINZ["black"])
        geometry = item["geometry"]
        narrative = geometry["narrative"]
        box(slide, narrative["x"], narrative["y"], narrative["width"], narrative["height"], AFINZ["black"])
        text_box(slide, item["narrative"], narrative["x"] + 12, narrative["y"] + 14,
                 narrative["width"] - 24, narrative["height"] - 28, 11, AFINZ["white"])

        if item["slide_instance_id"] == "c4":
            summary = item["summary"]
            delta = summary["delta"]
            text_box(slide, "Realizado", 28, 105, 120, 16, 9, AFINZ["black"])
            text_box(slide, compact_number(summary["current"]), 28, 120, 120, 34, 23, AFINZ["cyan"], True)
            text_box(slide, "Período equivalente", 170, 105, 140, 16, 9, AFINZ["black"])
            text_box(slide, compact_number(summary["previous"]), 170, 120, 140, 34, 23, AFINZ["black"], True)
            text_box(slide, "Variação", 330, 105, 115, 16, 9, AFINZ["black"])
            delta_text = "indisponível" if delta is None else f"{delta:+.1%}".replace(".", ",")
            text_box(slide, delta_text, 330, 120, 115, 34, 23, AFINZ["lime"], True)

            chart_data = ChartData()
            chart_data.categories = [str(value) for value in item["chart"]["categories"]]
            for series in item["chart"]["series"]:
                chart_data.add_series(series["name"], [value or 0 for value in series["values"]])
            visual = geometry["visual"]
            chart = slide.shapes.add_chart(
                XL_CHART_TYPE.LINE_MARKERS,
                int(visual["x"] * sx), int(visual["y"] * sy),
                int(visual["width"] * sx), int(visual["height"] * sy), chart_data,
            ).chart
            chart.has_title = False
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.value_axis.minimum_scale = 0
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.tick_labels.number_format = "#,##0"
            chart.value_axis.tick_labels.font.name = "Calibri"
            chart.value_axis.tick_labels.font.size = Pt(8)
            chart.category_axis.tick_labels.font.name = "Calibri"
            chart.category_axis.tick_labels.font.size = Pt(7)
            chart.legend.font.name = "Calibri"
            chart.legend.font.size = Pt(8)
            for series, contract in zip(chart.series, item["chart"]["series"]):
                series.format.line.color.rgb = rgb(contract["color"])
                series.format.line.width = Pt(2.25)
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.size = 4
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = rgb(contract["color"])
                series.marker.format.line.color.rgb = rgb(contract["color"])
        else:
            visual = geometry["visual"]
            metrics = item.get("metrics", [])
            column_width = visual["width"] / max(1, len(metrics))
            for index, metric in enumerate(metrics):
                x = visual["x"] + index * column_width
                if index:
                    box(slide, x, visual["y"] + 8, 1.5, visual["height"] - 16, AFINZ["black"])
                text_box(slide, metric["label"], x + 12, visual["y"] + 10,
                         column_width - 24, 20, 11, AFINZ["black"], True)
                text_box(slide, metric["value"], x + 12, visual["y"] + 34,
                         column_width - 24, 42, 30, AFINZ["cyan"], True)
                text_box(slide, metric["delta"], x + 12, visual["y"] + 82,
                         column_width - 24, 22, 13, AFINZ["black"], True)
                text_box(slide, metric["range"], x + 12, visual["y"] + 110,
                         column_width - 24, 18, 9, AFINZ["black"])
                text_box(slide, metric["verdict"], x + 12, visual["y"] + 130,
                         column_width - 24, 18, 10, AFINZ["lime"], True)
            if item.get("support_text"):
                text_box(slide, item["support_text"], 28, 350, 438, 18, 8, AFINZ["black"])

        slide.shapes.add_picture(str(logo), int(590 * sx), int(369 * sy), width=int(50 * sx))
        footer(slide, page)

    prs.core_properties.title = "Report Live - prova vertical Office"
    prs.core_properties.subject = f"Run certificado {plan['run_id']}"
    prs.core_properties.comments = "Gerado sem Google a partir do artefato imutável Report Live."
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def _wrap_pdf(text: str, font: str, size: float, width: float) -> list[str]:
    from reportlab.pdfbase.pdfmetrics import stringWidth

    safe = clean_text(text).replace("▲", "+").replace("▼", "-").replace("·", "/")
    lines: list[str] = []
    for paragraph in safe.splitlines():
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if stringWidth(candidate, font, size) <= width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def render_pdf(plan: dict[str, Any], logo: Path, output: Path) -> None:
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen.canvas import Canvas

    page_width, page_height = 960.0, 540.0
    sx, sy = page_width / CANVAS_WIDTH_PT, page_height / CANVAS_HEIGHT_PT
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(str(output), pagesize=(page_width, page_height), pageCompression=1)

    def color(name_or_hex: str) -> Any:
        value = AFINZ.get(name_or_hex, name_or_hex).lstrip("#")
        return HexColor(f"#{value}")

    def rect(x: float, y: float, w: float, h: float, fill: str) -> None:
        canvas.setFillColor(color(fill))
        canvas.setStrokeColor(color(fill))
        canvas.rect(x * sx, page_height - (y + h) * sy, w * sx, h * sy, fill=1, stroke=0)

    def text(text_value: str, x: float, y: float, w: float, h: float,
             size: float, fill: str, bold: bool = False) -> None:
        font = "Helvetica-Bold" if bold else "Helvetica"
        canvas.setFillColor(color(fill))
        canvas.setFont(font, size)
        line_height = size * 1.2
        cursor = page_height - y * sy - size
        for line in _wrap_pdf(text_value, font, size, w * sx):
            if cursor < page_height - (y + h) * sy:
                break
            canvas.drawString(x * sx, cursor, line)
            cursor -= line_height

    def footer(page: int, dark: bool = False) -> None:
        fill = "white" if dark else "black"
        text(f"Report Live. Run {plan['run_id'][:8]}. Spec {plan['versions']['spec']}.",
             28, 374, 500, 16, 8, fill)
        text(str(page), 670, 374, 22, 16, 8, fill)

    def line_chart(item: dict[str, Any]) -> None:
        geometry = item["geometry"]["visual"]
        x0, y0 = geometry["x"] * sx, page_height - (geometry["y"] + geometry["height"]) * sy
        width, height = geometry["width"] * sx, geometry["height"] * sy
        series = item["chart"]["series"]
        maximum = max(value for part in series for value in part["values"] if value is not None)
        maximum = max(1.0, maximum * 1.08)
        canvas.setStrokeColor(color("black"))
        canvas.setLineWidth(0.8)
        canvas.line(x0, y0, x0, y0 + height)
        canvas.line(x0, y0, x0 + width, y0)
        for tick in range(5):
            y = y0 + height * tick / 4
            canvas.setStrokeColor(color("black"))
            canvas.setLineWidth(0.25)
            canvas.line(x0, y, x0 + width, y)
            canvas.setFillColor(color("black"))
            canvas.setFont("Helvetica", 7)
            canvas.drawRightString(x0 - 5, y - 2, compact_number(maximum * tick / 4))
        categories = item["chart"]["categories"]
        for part in series:
            points = []
            for index, value in enumerate(part["values"]):
                if value is None:
                    continue
                x = x0 + width * index / max(1, len(categories) - 1)
                y = y0 + height * value / maximum
                points.append((x, y))
            canvas.setStrokeColor(color(part["color"]))
            canvas.setLineWidth(2.2)
            path = canvas.beginPath()
            for index, (x, y) in enumerate(points):
                (path.moveTo if index == 0 else path.lineTo)(x, y)
            canvas.drawPath(path, stroke=1, fill=0)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(color("black"))
        canvas.drawString(x0, y0 - 13, "Dia 1")
        canvas.drawRightString(x0 + width, y0 - 13, f"Dia {categories[-1]}")
        legend_y = y0 - 28
        for index, part in enumerate(series):
            lx = x0 + index * 130
            canvas.setStrokeColor(color(part["color"]))
            canvas.setLineWidth(2.2)
            canvas.line(lx, legend_y + 3, lx + 18, legend_y + 3)
            canvas.setFillColor(color("black"))
            canvas.setFont("Helvetica", 8)
            canvas.drawString(lx + 24, legend_y, part["name"])

    for page, item in enumerate(plan["slides"], start=1):
        if item["slide_instance_id"] == "c0":
            rect(0, 0, CANVAS_WIDTH_PT, CANVAS_HEIGHT_PT, "black")
            rect(28, 66, 76, 5, "cyan")
            rect(108, 66, 40, 5, "lime")
            text("report live", 28, 96, 500, 62, 34, "white", True)
            text(month_label(plan["period_start"]), 28, 158, 500, 34, 20, "cyan", True)
            text("prova vertical em PowerPoint e PDF", 28, 204, 500, 28, 14, "white")
            manifest = item["manifest"]
            source_cutoffs = json.loads(manifest.get("source_cutoffs", "{}"))
            contract = (
                f"Perfil {manifest.get('report_profile', '')}. Qualidade {manifest.get('quality_status', '')}. "
                f"Cutoff CRM {source_cutoffs.get('crm', 'indisponível')}."
            )
            text(contract, 28, 262, 560, 50, 12, "white")
            rect(600, 42, 80, 34, "white")
            canvas.drawImage(str(logo), 606 * sx, page_height - 72 * sy,
                             width=68 * sx, height=26 * sy,
                             preserveAspectRatio=True, mask="auto")
            footer(page, dark=True)
            canvas.showPage()
            continue

        rect(0, 0, CANVAS_WIDTH_PT, CANVAS_HEIGHT_PT, "white")
        rect(0, 0, CANVAS_WIDTH_PT, 8, "black")
        text(item["title"].lower(), 28, 28, 570, 38, 24, "black", True)
        text(f"Confiança {item['confidence']}. Fonte {item['source_view']}.", 28, 68, 570, 20, 9, "black")
        narrative = item["geometry"]["narrative"]
        rect(narrative["x"], narrative["y"], narrative["width"], narrative["height"], "black")
        text(item["narrative"], narrative["x"] + 12, narrative["y"] + 14,
             narrative["width"] - 24, narrative["height"] - 28, 11, "white")

        if item["slide_instance_id"] == "c4":
            summary = item["summary"]
            text("Realizado", 28, 105, 120, 16, 9, "black")
            text(compact_number(summary["current"]), 28, 120, 120, 34, 23, "cyan", True)
            text("Período equivalente", 170, 105, 140, 16, 9, "black")
            text(compact_number(summary["previous"]), 170, 120, 140, 34, 23, "black", True)
            delta = summary["delta"]
            delta_text = "indisponível" if delta is None else f"{delta:+.1%}".replace(".", ",")
            text("Variação", 330, 105, 115, 16, 9, "black")
            text(delta_text, 330, 120, 115, 34, 23, "lime", True)
            line_chart(item)
        else:
            visual = item["geometry"]["visual"]
            metrics = item.get("metrics", [])
            width = visual["width"] / max(1, len(metrics))
            for index, metric in enumerate(metrics):
                x = visual["x"] + index * width
                if index:
                    rect(x, visual["y"] + 8, 1.5, visual["height"] - 16, "black")
                text(metric["label"], x + 12, visual["y"] + 10, width - 24, 20, 11, "black", True)
                text(metric["value"], x + 12, visual["y"] + 34, width - 24, 42, 30, "cyan", True)
                text(metric["delta"], x + 12, visual["y"] + 82, width - 24, 22, 13, "black", True)
                text(metric["range"], x + 12, visual["y"] + 110, width - 24, 18, 9, "black")
                text(metric["verdict"], x + 12, visual["y"] + 130, width - 24, 18, 10, "lime", True)
            if item.get("support_text"):
                text(item["support_text"], 28, 350, 438, 18, 8, "black")

        canvas.drawImage(str(logo), 590 * sx, page_height - 392 * sy,
                         width=50 * sx, height=20 * sy,
                         preserveAspectRatio=True, mask="auto")
        footer(page)
        canvas.showPage()

    canvas.save()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifact", type=Path, required=True)
    parser.add_argument("--expected-run-id", required=True)
    parser.add_argument("--expected-content-hash", required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--logo", type=Path, required=True)
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--plan", type=Path)
    parser.add_argument("--slides", default=",".join(DEFAULT_SLIDES))
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    artifact = json.loads(args.artifact.read_text(encoding="utf-8"))
    validate_artifact_identity(artifact, args.expected_run_id, args.expected_content_hash)
    slide_ids = tuple(value.strip() for value in args.slides.split(",") if value.strip())
    plan = build_render_plan(artifact, slide_ids)
    if args.plan:
        args.plan.parent.mkdir(parents=True, exist_ok=True)
        args.plan.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    render_pptx(plan, args.template, args.logo, args.pptx)
    render_pdf(plan, args.logo, args.pdf)
    print(json.dumps({
        "run_id": plan["run_id"],
        "slides": [item["slide_instance_id"] for item in plan["slides"]],
        "pptx": str(args.pptx.resolve()),
        "pdf": str(args.pdf.resolve()),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
