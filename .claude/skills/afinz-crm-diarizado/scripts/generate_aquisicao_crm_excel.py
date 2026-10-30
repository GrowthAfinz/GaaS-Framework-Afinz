#!/usr/bin/env python3
"""Generate the AFINZ CRM acquisition daily Excel report."""

from __future__ import annotations

import argparse
import json
import os
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_SUPABASE_URL = "https://mipiwxadnpwtcgfcedym.supabase.co"
DEFAULT_SUPABASE_KEY = os.environ.get("SUPABASE_ANON_KEY", "")
DEFAULT_START = date(2026, 4, 13)
DEFAULT_END = date(2026, 5, 21)

SECTIONS: list[tuple[str, list[str]]] = [
    ("B2C", ["TOPO DE FUNIL B2C", "REPESCAGEM B2C", "UPGRADE B2C", "LEADS PARCEIROS B2C", "CARRINHO ABANDONADO B2C"]),
    ("PLURIX", ["TOPO DE FUNIL PLURIX", "REPESCAGEM PLURIX", "UPGRADE PLURIX", "RECENCIA PLURIX", "CARRINHO ABANDONADO PLURIX"]),
    ("DIA", ["TOPO DE FUNIL DIA", "REPESCAGEM DIA"]),
    ("BEM BARATO", ["TOPO DE FUNIL BB", "REPESCAGEM BB"]),
]

CHANNEL_ORDER = {"E-MAIL": 0, "SMS": 1, "WPP": 2}
NUMERIC_FIELDS = ("entregues", "emitidos", "aprovados", "propostas")
B2C_REALIZED_FIELDS = ("propostas_total", "emissoes_total", "propostas_serasa", "emissoes_serasa")
DAY_NAMES = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sab", "Dom"]

COLORS = {
    "header_gray": "D9D9D9",
    "section_dark": "17365D",
    "data_b2c": "1F4E79",
    "data_b2c_sub": "2F5496",
    "topo_b2c": "1F6B3A",
    "topo": "1F4E79",
    "repescagem": "2E7D32",
    "upgrade": "B8860B",
    "leads": "00695C",
    "recencia": "6A1B9A",
    "email_fill": "DCE6F1",
    "email_font": "1F4E79",
    "wpp_fill": "E2EFDA",
    "wpp_font": "375623",
    "sms_fill": "FFF2CC",
    "sms_font": "7F6000",
    "zebra": "F5F5F5",
    "weekend": "E8E8E8",
    "manual_b2c": "EBF3FF",
    "manual_b2c_repeat": "B4C7E7",
    "current_period": "E2F0F9",
    "total": "D9EAD3",
    "audit_header": "1F4E79",
    "audit_subheader": "D9EAF7",
}


def parse_iso_date(value: str) -> date:
    return datetime.strptime(value, "%Y-%m-%d").date()


def parse_row_date(value: Any) -> date:
    if not value:
        raise ValueError("missing Data de Disparo")
    text = str(value)
    if "T" in text:
        text = text.split("T", 1)[0]
    return parse_iso_date(text[:10])


def as_int(value: Any) -> int:
    if value in (None, ""):
        return 0
    return int(float(value))


def normalize_channel(value: Any) -> str:
    raw = str(value or "").upper()
    if "MAIL" in raw:
        return "E-MAIL"
    if "WHATSAPP" in raw or raw == "WPP":
        return "WPP"
    if "SMS" in raw:
        return "SMS"
    return raw.strip() or "N/A"


def classify(row: dict[str, Any]) -> tuple[str, str] | None:
    bu = row.get("BU")
    parceiro = row.get("Parceiro")
    segmento = row.get("Segmento")
    etapa = row.get("Etapa de aquisição")

    if segmento == "Leads_Parceiros":
        return "B2C", "LEADS PARCEIROS B2C"
    if bu == "B2C" and parceiro == "Serasa" and segmento == "Base_Proprietaria" and etapa == "Meio_de_Funil":
        return "B2C", "LEADS PARCEIROS B2C"
    if bu == "B2C" and parceiro == "Proprietaria" and segmento == "Base_Proprietaria" and etapa == "Aquisicao":
        return "B2C", "TOPO DE FUNIL B2C"
    if bu == "B2C" and parceiro == "Proprietaria" and segmento == "Negados" and etapa == "Meio_de_Funil":
        return "B2C", "REPESCAGEM B2C"
    if bu == "B2C" and parceiro == "Proprietaria" and segmento == "Aprovados_nao_convertidos" and etapa == "Meio_de_Funil":
        return "B2C", "UPGRADE B2C"
    if bu == "B2C" and parceiro in {"N/A", "Serasa"} and segmento == "Abandonados":
        return "B2C", "CARRINHO ABANDONADO B2C"

    if bu == "Plurix" and parceiro == "N/A" and segmento == "CRM" and etapa == "Aquisicao":
        return "PLURIX", "TOPO DE FUNIL PLURIX"
    if bu == "Plurix" and parceiro == "N/A" and segmento == "Negados" and etapa == "Meio_de_Funil":
        return "PLURIX", "REPESCAGEM PLURIX"
    if bu == "Plurix" and parceiro == "N/A" and segmento == "Aprovados_nao_convertidos" and etapa == "Meio_de_Funil":
        return "PLURIX", "UPGRADE PLURIX"
    if bu == "Plurix" and parceiro == "N/A" and segmento == "Abandonados":
        return "PLURIX", "CARRINHO ABANDONADO PLURIX"
    if bu == "Plurix" and parceiro == "N/A" and segmento == "Recencia de Compra":
        return "PLURIX", "RECENCIA PLURIX"

    if bu == "B2B2C" and parceiro == "Dia" and etapa == "Aquisicao":
        return "DIA", "TOPO DE FUNIL DIA"
    if bu == "B2B2C" and parceiro == "Dia" and etapa == "Meio_de_Funil":
        return "DIA", "REPESCAGEM DIA"

    if bu == "B2B2C" and parceiro == "Bem Barato" and segmento in {"CRM", "Recencia de Compra"} and etapa == "Aquisicao":
        return "BEM BARATO", "TOPO DE FUNIL BB"
    if (
        bu == "B2B2C"
        and parceiro == "Bem Barato"
        and segmento in {"Negados", "Aprovados_nao_convertidos"}
        and etapa == "Meio_de_Funil"
    ):
        return "BEM BARATO", "REPESCAGEM BB"
    return None


def audit_reason(row: dict[str, Any]) -> str:
    if row.get("BU") == "Seguros":
        return "Fora do escopo: Seguros/Rentabilizacao"
    if row.get("Segmento") == "Abandonados":
        return "Sem bloco definido: Abandonados"
    return "Sem regra de mapeamento"


def fetch_supabase_rows(supabase_url: str, supabase_key: str, start: date, end: date) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    page_size = 1000
    endpoint = supabase_url.rstrip("/") + "/rest/v1/activities"
    exclusive_end = end + timedelta(days=1)

    while True:
        params = urlencode(
            {
                "select": "*",
                "Data de Disparo": f"gte.{start.isoformat()}",
                "order": "Data de Disparo.asc,BU.asc,Parceiro.asc",
                "limit": str(page_size),
                "offset": str(offset),
            }
        )
        # urllib cannot encode two filters with the same key through a dict.
        params += "&" + urlencode({"Data de Disparo": f"lt.{exclusive_end.isoformat()}"})
        request = Request(
            endpoint + "?" + params,
            headers={
                "apikey": supabase_key,
                "Authorization": f"Bearer {supabase_key}",
                "Accept": "application/json",
            },
        )
        try:
            with urlopen(request, timeout=45) as response:
                page = json.loads(response.read().decode("utf-8"))
        except HTTPError as exc:
            message = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"Supabase request failed ({exc.code}): {message}") from exc

        rows.extend(page)
        if len(page) < page_size:
            break
        offset += page_size

    return rows


def fetch_b2c_daily_metrics(supabase_url: str, supabase_key: str, start: date, end: date) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    page_size = 1000
    endpoint = supabase_url.rstrip("/") + "/rest/v1/b2c_daily_metrics"

    while True:
        params = urlencode(
            {
                "select": "*",
                "data": f"gte.{start.isoformat()}",
                "order": "data.asc,tipo.asc",
                "limit": str(page_size),
                "offset": str(offset),
            }
        )
        params += "&" + urlencode({"data": f"lte.{end.isoformat()}"})
        request = Request(
            endpoint + "?" + params,
            headers={
                "apikey": supabase_key,
                "Authorization": f"Bearer {supabase_key}",
                "Accept": "application/json",
            },
        )
        try:
            with urlopen(request, timeout=45) as response:
                page = json.loads(response.read().decode("utf-8"))
        except HTTPError as exc:
            message = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"Supabase b2c_daily_metrics request failed ({exc.code}): {message}") from exc

        rows.extend(page)
        if len(page) < page_size:
            break
        offset += page_size

    return rows


def build_b2c_realized_index(
    rows: list[dict[str, Any]],
    start: date,
    end: date,
    overrides: list[dict[str, Any]] | None = None,
) -> dict[str, dict[str, int]]:
    idx: dict[str, dict[str, int]] = defaultdict(
        lambda: {
            "propostas_total": 0,
            "emissoes_total": 0,
            "propostas_serasa": 0,
            "emissoes_serasa": 0,
        }
    )

    for row in rows:
        if not row.get("data"):
            continue
        row_date = parse_iso_date(str(row.get("data"))[:10])
        if row_date < start or row_date > end:
            continue

        tipo = str(row.get("tipo") or "").strip().lower()
        if tipo == "total":
            idx[row_date.isoformat()]["propostas_total"] += as_int(row.get("propostas_total"))
            idx[row_date.isoformat()]["emissoes_total"] += as_int(row.get("emissoes_total"))
        elif "serasa" in tipo:
            idx[row_date.isoformat()]["propostas_serasa"] += as_int(row.get("propostas_total"))
            idx[row_date.isoformat()]["emissoes_serasa"] += as_int(row.get("emissoes_total"))

    if overrides:
        for row in overrides:
            raw_date = row.get("data") or row.get("date") or row.get("Date")
            if not raw_date:
                continue
            row_date = parse_iso_date(str(raw_date)[:10])
            if row_date < start or row_date > end:
                continue
            ds = row_date.isoformat()
            for field in B2C_REALIZED_FIELDS:
                if field in row:
                    idx[ds][field] = as_int(row.get(field))
            if "propostas_total" not in row and ("propostas_app" in row or "propostas_serasa" in row):
                idx[ds]["propostas_total"] = as_int(row.get("propostas_app")) + as_int(row.get("propostas_serasa"))
            if "emissoes_total" not in row and ("emissoes_app" in row or "emissoes_serasa" in row):
                idx[ds]["emissoes_total"] = as_int(row.get("emissoes_app")) + as_int(row.get("emissoes_serasa"))

    return {ds: {field: values[field] for field in B2C_REALIZED_FIELDS} for ds, values in idx.items()}


def build_indexes(
    rows: list[dict[str, Any]], start: date, end: date
) -> tuple[
    dict[tuple[str, str, str, str], dict[str, int]],
    list[dict[str, Any]],
    list[dict[str, Any]],
    dict[str, int],
]:
    idx: dict[tuple[str, str, str, str], dict[str, int]] = defaultdict(lambda: {"entregues": 0, "emitidos": 0, "aprovados": 0, "propostas": 0, "linhas_origem": 0})
    audit: dict[tuple[str, str, str, str, str, str, str], dict[str, Any]] = {}
    journey_map: dict[tuple[str, str, str, str, str, str, str, str, str, str], dict[str, Any]] = {}
    summary = {"source_rows": 0, "mapped_source_rows": 0, "audit_source_rows": 0}

    for row in rows:
        row_date = parse_row_date(row.get("Data de Disparo"))
        if row_date < start or row_date > end:
            continue
        summary["source_rows"] += 1
        canal = normalize_channel(row.get("Canal"))
        metrics = {
            "entregues": as_int(row.get("Base Acionável")),
            "emitidos": as_int(row.get("Cartões Gerados")),
            "aprovados": as_int(row.get("Aprovados")),
            "propostas": as_int(row.get("Propostas")),
        }

        mapped = classify(row)
        destino = "MAPEADO" if mapped else audit_reason(row)
        secao = mapped[0] if mapped else ""
        bloco = mapped[1] if mapped else ""
        journey_key = (
            destino,
            secao,
            bloco,
            str(row.get("BU") or ""),
            str(row.get("Parceiro") or ""),
            str(row.get("Segmento") or ""),
            str(row.get("Etapa de aquisição") or ""),
            canal,
            str(row.get("jornada") or ""),
            str(row.get("Activity name / Taxonomia") or ""),
        )
        if journey_key not in journey_map:
            journey_map[journey_key] = {
                "status_mapeamento": destino,
                "secao": secao,
                "bloco": bloco,
                "bu": journey_key[3],
                "parceiro": journey_key[4],
                "segmento": journey_key[5],
                "etapa": journey_key[6],
                "canal": canal,
                "jornada": journey_key[8],
                "taxonomia": journey_key[9],
                "linhas": 0,
                "entregues": 0,
                "emitidos": 0,
                "aprovados": 0,
                "propostas": 0,
            }
        journey_map[journey_key]["linhas"] += 1
        for field, value in metrics.items():
            journey_map[journey_key][field] += value

        if mapped:
            secao, bloco = mapped
            key = (row_date.isoformat(), secao, bloco, canal)
            for field, value in metrics.items():
                idx[key][field] += value
            idx[key]["linhas_origem"] += 1
            summary["mapped_source_rows"] += 1
        else:
            audit_key = (
                audit_reason(row),
                str(row.get("BU") or ""),
                str(row.get("Parceiro") or ""),
                str(row.get("Segmento") or ""),
                str(row.get("Etapa de aquisição") or ""),
                str(row.get("Canal") or ""),
                canal,
            )
            if audit_key not in audit:
                audit[audit_key] = {
                    "motivo": audit_key[0],
                    "bu": audit_key[1],
                    "parceiro": audit_key[2],
                    "segmento": audit_key[3],
                    "etapa": audit_key[4],
                    "canal_raw": audit_key[5],
                    "canal": audit_key[6],
                    "linhas": 0,
                    "entregues": 0,
                    "emitidos": 0,
                    "aprovados": 0,
                    "propostas": 0,
                }
            audit[audit_key]["linhas"] += 1
            for field, value in metrics.items():
                audit[audit_key][field] += value
            summary["audit_source_rows"] += 1

    return (
        idx,
        sorted(audit.values(), key=lambda r: (r["motivo"], r["bu"], r["parceiro"], r["segmento"], r["etapa"], r["canal"])),
        sorted(
            journey_map.values(),
            key=lambda r: (
                r["status_mapeamento"],
                r["secao"],
                r["bloco"],
                r["bu"],
                r["parceiro"],
                r["segmento"],
                r["etapa"],
                r["canal"],
                r["jornada"],
                r["taxonomia"],
            ),
        ),
        summary,
    )


def all_dates(start: date, end: date) -> list[date]:
    return [start + timedelta(days=i) for i in range((end - start).days + 1)]


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def block_color(section: str, block: str) -> str:
    if block.startswith("TOPO"):
        return COLORS["topo_b2c"] if section == "B2C" else COLORS["topo"]
    if block.startswith("REPESCAGEM"):
        return COLORS["repescagem"]
    if block.startswith("UPGRADE"):
        return COLORS["upgrade"]
    if block.startswith("LEADS"):
        return COLORS["leads"]
    if block.startswith("CARRINHO"):
        return COLORS["leads"]
    if block.startswith("RECENCIA"):
        return COLORS["recencia"]
    return COLORS["section_dark"]


def channel_style(canal: str) -> tuple[str, str]:
    if canal == "E-MAIL":
        return COLORS["email_fill"], COLORS["email_font"]
    if canal == "WPP":
        return COLORS["wpp_fill"], COLORS["wpp_font"]
    if canal == "SMS":
        return COLORS["sms_fill"], COLORS["sms_font"]
    return "FFFFFF", "000000"


def set_cell(cell, value=None, *, bold=False, italic=False, font_color="000000", fill_color=None, align="center"):
    if value is not None:
        cell.value = value
    cell.font = Font(name="Calibri", size=10, bold=bold, italic=italic, color=font_color)
    if fill_color:
        cell.fill = fill(fill_color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)


def apply_border(ws, row: int, max_col: int, top_style: str = "thin", vertical_boundaries: list[int] | None = None) -> None:
    vertical_boundaries = vertical_boundaries or []
    top = Side(style=top_style, color="BFBFBF")
    thin = Side(style="thin", color="E0E0E0")
    divider = Side(style="medium", color="000000")
    for col in range(1, max_col + 1):
        ws.cell(row, col).border = Border(
            top=top,
            bottom=thin,
            left=thin,
            right=divider if col in vertical_boundaries else thin,
        )


def apply_merged_header_divider(ws, row: int, start_col: int, top_style: str = "medium") -> None:
    cell = ws.cell(row, start_col)
    top = Side(style=top_style, color="BFBFBF")
    thin = Side(style="thin", color="E0E0E0")
    divider = Side(style="medium", color="000000")
    cell.border = Border(top=top, bottom=thin, left=thin, right=divider)


def apply_b2c_repeat_fill(ws, row: int) -> None:
    divider = Side(style="medium", color="000000")
    empty = Side(style=None)
    for col in range(3, 7):
        ws.cell(row, col).fill = fill(COLORS["manual_b2c_repeat"])
        ws.cell(row, col).border = Border(
            left=empty,
            right=divider if col == 6 else empty,
            top=empty,
            bottom=empty,
        )


def write_section(
    ws,
    start_row: int,
    section: str,
    blocks: list[str],
    idx: dict[tuple[str, str, str, str], dict[str, int]],
    dates: list[date],
    b2c_realized_idx: dict[str, dict[str, int]],
) -> int:
    ws.cell(start_row, 1, "Data")
    ws.cell(start_row, 2, "Dia")
    for col in (1, 2):
        set_cell(ws.cell(start_row, col), bold=True, fill_color=COLORS["header_gray"])
        set_cell(ws.cell(start_row + 1, col), fill_color=COLORS["header_gray"])

    col = 3
    numeric_cols_by_section: list[int] = []
    vertical_boundaries = [2]

    if section == "B2C":
        ws.merge_cells(start_row=start_row, start_column=col, end_row=start_row, end_column=col + 3)
        set_cell(ws.cell(start_row, col), "DADOS REALIZADOS B2C", bold=True, font_color="FFFFFF", fill_color=COLORS["data_b2c"])
        merged_header_starts = [col]
        for offset, label in enumerate(["Propostas", "Cartões Emitidos", "Propostas Serasa", "Cartões Serasa"]):
            set_cell(ws.cell(start_row + 1, col + offset), label, bold=True, font_color="FFFFFF", fill_color=COLORS["data_b2c_sub"])
        numeric_cols_by_section.extend([col, col + 1, col + 2, col + 3])
        vertical_boundaries.append(col + 3)
        col += 4
    else:
        merged_header_starts = []

    block_positions: dict[str, int] = {}
    for block in blocks:
        block_positions[block] = col
        ws.merge_cells(start_row=start_row, start_column=col, end_row=start_row, end_column=col + 4)
        merged_header_starts.append(col)
        color = block_color(section, block)
        set_cell(ws.cell(start_row, col), block, bold=True, font_color="FFFFFF", fill_color=color)
        for offset, label in enumerate(["Canal", "Entregues", "Emitidos", "Aprovados", "Proposta"]):
            set_cell(ws.cell(start_row + 1, col + offset), label, bold=True, font_color="FFFFFF", fill_color=color)
        numeric_cols_by_section.extend([col + 1, col + 2, col + 3, col + 4])
        vertical_boundaries.append(col + 4)
        col += 5

    max_col = col - 1
    apply_border(ws, start_row, max_col, "medium", vertical_boundaries)
    apply_border(ws, start_row + 1, max_col, "thin", vertical_boundaries)
    for header_start_col in merged_header_starts:
        apply_merged_header_divider(ws, start_row, header_start_col)
    data_start = start_row + 2
    row_num = data_start

    for i, day in enumerate(dates):
        ds = day.isoformat()
        canais = set()
        for block in blocks:
            for key in idx:
                rd, secao, rb, canal = key
                if rd == ds and secao == section and rb == block:
                    canais.add(canal)
        rows_to_write = sorted(canais, key=lambda c: (CHANNEL_ORDER.get(c, 99), c)) or [None]
        is_current_or_future_month_day = day >= date.today() and day.month == date.today().month and day.year == date.today().year
        zebra = COLORS["current_period"] if is_current_or_future_month_day else (COLORS["zebra"] if i % 2 else "FFFFFF")

        for channel_index, canal in enumerate(rows_to_write):
            for c in range(1, max_col + 1):
                set_cell(ws.cell(row_num, c), fill_color=zebra)

            date_fill = COLORS["weekend"] if day.weekday() >= 5 else zebra
            set_cell(ws.cell(row_num, 1), day.strftime("%d/%m/%Y"), italic=day.weekday() >= 5, fill_color=date_fill)
            set_cell(ws.cell(row_num, 2), DAY_NAMES[day.weekday()], italic=day.weekday() >= 5, fill_color=date_fill)

            if section == "B2C":
                for c in range(3, 7):
                    set_cell(
                        ws.cell(row_num, c),
                        fill_color=COLORS["manual_b2c"] if channel_index == 0 else COLORS["manual_b2c_repeat"],
                    )
                if channel_index == 0:
                    realized = b2c_realized_idx.get(ds)
                    if realized:
                        for offset, field in enumerate(B2C_REALIZED_FIELDS, start=3):
                            value = realized[field]
                            if value:
                                set_cell(ws.cell(row_num, offset), value, fill_color=COLORS["manual_b2c"])

            if canal:
                for block in blocks:
                    values = idx.get((ds, section, block, canal))
                    if values and (values["entregues"] > 0 or values["emitidos"] > 0 or values["aprovados"] > 0 or values["propostas"] > 0):
                        start_col = block_positions[block]
                        channel_fill, font_color = channel_style(canal)
                        set_cell(ws.cell(row_num, start_col), canal, bold=True, font_color=font_color, fill_color=channel_fill)
                        for offset, field in enumerate(NUMERIC_FIELDS, start=1):
                            set_cell(ws.cell(row_num, start_col + offset), values[field], font_color=font_color, fill_color=channel_fill)

            apply_border(ws, row_num, max_col, "medium" if channel_index == 0 else "thin", vertical_boundaries)
            if section == "B2C" and channel_index > 0:
                apply_b2c_repeat_fill(ws, row_num)
            row_num += 1

    total_row = row_num
    for c in range(1, max_col + 1):
        set_cell(ws.cell(total_row, c), fill_color=COLORS["total"], bold=True)
    ws.cell(total_row, 1, "total")
    for c in numeric_cols_by_section:
        letter = get_column_letter(c)
        ws.cell(total_row, c, f"=SUM({letter}{data_start}:{letter}{total_row - 1})")
    apply_border(ws, total_row, max_col, "medium", vertical_boundaries)
    return total_row + 3


def write_audit_sheet(wb: Workbook, audit_rows: list[dict[str, Any]], journey_rows: list[dict[str, Any]], summary: dict[str, int]) -> None:
    ws = wb.create_sheet("Auditoria")
    ws.freeze_panes = "A9"

    set_cell(ws["A1"], "Auditoria de mapeamento", bold=True, font_color="FFFFFF", fill_color=COLORS["audit_header"], align="left")
    ws.merge_cells("A1:K1")

    summary_rows = [
        ("Linhas fonte no período", summary["source_rows"]),
        ("Linhas mapeadas no relatório", summary["mapped_source_rows"]),
        ("Linhas fora do relatório", summary["audit_source_rows"]),
    ]
    for idx_row, (label, value) in enumerate(summary_rows, start=3):
        set_cell(ws.cell(idx_row, 1), label, bold=True, fill_color=COLORS["audit_subheader"], align="left")
        set_cell(ws.cell(idx_row, 2), value, fill_color=COLORS["audit_subheader"])

    reason_totals: dict[str, dict[str, int]] = defaultdict(lambda: {"linhas": 0, "entregues": 0, "emitidos": 0, "aprovados": 0, "propostas": 0})
    for row in audit_rows:
        reason = row["motivo"]
        for field in ("linhas", "entregues", "emitidos", "aprovados", "propostas"):
            reason_totals[reason][field] += int(row[field])

    headers = ["Motivo", "Linhas", "Entregues", "Emitidos", "Aprovados", "Propostas"]
    for col, label in enumerate(headers, start=4):
        set_cell(ws.cell(3, col), label, bold=True, font_color="FFFFFF", fill_color=COLORS["audit_header"])
    reason_row = 4
    for reason, totals in sorted(reason_totals.items()):
        set_cell(ws.cell(reason_row, 4), reason, align="left")
        for offset, field in enumerate(("linhas", "entregues", "emitidos", "aprovados", "propostas"), start=5):
            set_cell(ws.cell(reason_row, offset), totals[field])
        reason_row += 1

    mapped_totals: dict[tuple[str, str], dict[str, int]] = defaultdict(
        lambda: {"linhas": 0, "entregues": 0, "emitidos": 0, "aprovados": 0, "propostas": 0}
    )
    for journey in journey_rows:
        if journey["status_mapeamento"] != "MAPEADO":
            continue
        key = (journey["secao"], journey["bloco"])
        for field in ("linhas", "entregues", "emitidos", "aprovados", "propostas"):
            mapped_totals[key][field] += int(journey[field])

    mapped_start = max(reason_row + 2, 8)
    set_cell(
        ws.cell(mapped_start, 1),
        "Resumo dos blocos mapeados",
        bold=True,
        font_color="FFFFFF",
        fill_color=COLORS["audit_header"],
        align="left",
    )
    ws.merge_cells(start_row=mapped_start, start_column=1, end_row=mapped_start, end_column=7)
    mapped_headers = ["Seção", "Bloco", "Linhas", "Entregues", "Emitidos", "Aprovados", "Propostas"]
    for col, label in enumerate(mapped_headers, start=1):
        set_cell(ws.cell(mapped_start + 1, col), label, bold=True, font_color="FFFFFF", fill_color=COLORS["audit_header"])
    mapped_row = mapped_start + 2
    for (secao, bloco), totals in sorted(mapped_totals.items()):
        set_cell(ws.cell(mapped_row, 1), secao, align="left")
        set_cell(ws.cell(mapped_row, 2), bloco, align="left")
        for offset, field in enumerate(("linhas", "entregues", "emitidos", "aprovados", "propostas"), start=3):
            set_cell(ws.cell(mapped_row, offset), totals[field])
        mapped_row += 1

    detail_start = mapped_row + 2
    detail_headers = ["Motivo", "BU", "Parceiro", "Segmento", "Etapa", "Canal original", "Canal", "Linhas", "Entregues", "Emitidos", "Aprovados", "Propostas"]
    for col, label in enumerate(detail_headers, start=1):
        set_cell(ws.cell(detail_start, col), label, bold=True, font_color="FFFFFF", fill_color=COLORS["audit_header"])

    for r, audit in enumerate(audit_rows, start=detail_start + 1):
        values = [
            audit["motivo"],
            audit["bu"],
            audit["parceiro"],
            audit["segmento"],
            audit["etapa"],
            audit["canal_raw"],
            audit["canal"],
            audit["linhas"],
            audit["entregues"],
            audit["emitidos"],
            audit["aprovados"],
            audit["propostas"],
        ]
        for c, value in enumerate(values, start=1):
            set_cell(ws.cell(r, c), value, align="left" if c <= 6 else "center")

    journey_start = detail_start + len(audit_rows) + 4
    set_cell(
        ws.cell(journey_start, 1),
        "Mapa de jornadas e taxonomias por destino",
        bold=True,
        font_color="FFFFFF",
        fill_color=COLORS["audit_header"],
        align="left",
    )
    ws.merge_cells(start_row=journey_start, start_column=1, end_row=journey_start, end_column=15)
    journey_headers = [
        "Status/Motivo",
        "Seção",
        "Bloco",
        "BU",
        "Parceiro",
        "Segmento",
        "Etapa",
        "Canal",
        "Jornada",
        "Activity name / Taxonomia",
        "Linhas",
        "Entregues",
        "Emitidos",
        "Aprovados",
        "Propostas",
    ]
    for col, label in enumerate(journey_headers, start=1):
        set_cell(ws.cell(journey_start + 1, col), label, bold=True, font_color="FFFFFF", fill_color=COLORS["audit_header"])

    for r, journey in enumerate(journey_rows, start=journey_start + 2):
        values = [
            journey["status_mapeamento"],
            journey["secao"],
            journey["bloco"],
            journey["bu"],
            journey["parceiro"],
            journey["segmento"],
            journey["etapa"],
            journey["canal"],
            journey["jornada"],
            journey["taxonomia"],
            journey["linhas"],
            journey["entregues"],
            journey["emitidos"],
            journey["aprovados"],
            journey["propostas"],
        ]
        for c, value in enumerate(values, start=1):
            set_cell(ws.cell(r, c), value, align="left" if c <= 10 else "center")

    widths = [48, 12, 18, 14, 18, 28, 22, 12, 42, 58, 10, 14, 12, 12, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_workbook(
    raw_rows: list[dict[str, Any]],
    start: date,
    end: date,
    b2c_daily_rows: list[dict[str, Any]] | None = None,
    b2c_realized_overrides: list[dict[str, Any]] | None = None,
) -> Workbook:
    idx, audit_rows, journey_rows, summary = build_indexes(raw_rows, start, end)
    b2c_realized_idx = build_b2c_realized_index(b2c_daily_rows or [], start, end, b2c_realized_overrides)
    dates = all_dates(start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "Aquisição CRM"
    ws.freeze_panes = "C3"
    ws.sheet_view.showGridLines = False

    next_row = 1
    for section, blocks in SECTIONS:
        next_row = write_section(ws, next_row, section, blocks, idx, dates, b2c_realized_idx)

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 6
    for c in range(3, ws.max_column + 1):
        letter = get_column_letter(c)
        header = str(ws.cell(2, c).value or "")
        if header == "Canal":
            ws.column_dimensions[letter].width = 10
        elif header in {"Cartões Emitidos", "Propostas Serasa", "Cartões Serasa"}:
            ws.column_dimensions[letter].width = 16
        else:
            ws.column_dimensions[letter].width = 12

    write_audit_sheet(wb, audit_rows, journey_rows, summary)
    return wb


def validate_workbook(path: Path, start: date, end: date) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    expected = {"Aquisição CRM", "Auditoria"}
    if set(wb.sheetnames) != expected:
        raise AssertionError(f"unexpected sheets: {wb.sheetnames}")
    ws = wb["Aquisição CRM"]
    if ws.freeze_panes != "C3":
        raise AssertionError(f"unexpected freeze panes: {ws.freeze_panes}")

    dates_found = []
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str) and len(value) == 10 and value[2] == "/" and value[5] == "/":
            dates_found.append(value)

    if start.strftime("%d/%m/%Y") not in dates_found:
        raise AssertionError("start date missing from report")
    if end.strftime("%d/%m/%Y") not in dates_found:
        raise AssertionError("end date missing from report")
    if any("/2024" in d or "/2025" in d or "/2027" in d or "/2028" in d or "/2029" in d for d in dates_found):
        raise AssertionError("invalid year copied into report dates")

    bad_terms = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(term in cell.value for term in bad_terms):
                    raise AssertionError(f"formula error marker in {sheet.title}!{cell.coordinate}: {cell.value}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate AFINZ CRM daily acquisition Excel report.")
    parser.add_argument("--start", default=DEFAULT_START.isoformat())
    parser.add_argument("--end", default=DEFAULT_END.isoformat())
    parser.add_argument("--output", default="outputs/aquisicao_crm_copa2026.xlsx")
    parser.add_argument("--supabase-url", default=os.environ.get("SUPABASE_URL", DEFAULT_SUPABASE_URL))
    parser.add_argument("--supabase-key", default=os.environ.get("SUPABASE_ANON_KEY", DEFAULT_SUPABASE_KEY))
    parser.add_argument("--raw-json", help="Optional local raw activities JSON file. If omitted, fetches Supabase REST.")
    parser.add_argument("--write-raw-json", help="Optional path to save fetched raw rows for audit/replay.")
    parser.add_argument("--b2c-raw-json", help="Optional local b2c_daily_metrics JSON file. If omitted, fetches Supabase REST.")
    parser.add_argument("--write-b2c-raw-json", help="Optional path to save fetched b2c_daily_metrics rows for audit/replay.")
    parser.add_argument("--b2c-realized-json", help="Optional realized B2C override JSON with data/date, propostas_total, emissoes_total, propostas_serasa, emissoes_serasa. Legacy propostas_app/emissoes_app are accepted and summed with Serasa.")
    args = parser.parse_args()

    start = parse_iso_date(args.start)
    end = parse_iso_date(args.end)

    if args.raw_json:
        with open(args.raw_json, "r", encoding="utf-8") as handle:
            raw_rows = json.load(handle)
    else:
        raw_rows = fetch_supabase_rows(args.supabase_url, args.supabase_key, start, end)

    if args.b2c_raw_json:
        with open(args.b2c_raw_json, "r", encoding="utf-8") as handle:
            b2c_daily_rows = json.load(handle)
    else:
        b2c_daily_rows = fetch_b2c_daily_metrics(args.supabase_url, args.supabase_key, start, end)

    if args.b2c_realized_json:
        with open(args.b2c_realized_json, "r", encoding="utf-8") as handle:
            b2c_realized_overrides = json.load(handle)
    else:
        b2c_realized_overrides = None

    if args.write_raw_json:
        raw_path = Path(args.write_raw_json)
        raw_path.parent.mkdir(parents=True, exist_ok=True)
        with raw_path.open("w", encoding="utf-8") as handle:
            json.dump(raw_rows, handle, ensure_ascii=False, indent=2)

    if args.write_b2c_raw_json:
        b2c_raw_path = Path(args.write_b2c_raw_json)
        b2c_raw_path.parent.mkdir(parents=True, exist_ok=True)
        with b2c_raw_path.open("w", encoding="utf-8") as handle:
            json.dump(b2c_daily_rows, handle, ensure_ascii=False, indent=2)

    workbook = build_workbook(raw_rows, start, end, b2c_daily_rows, b2c_realized_overrides)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    validate_workbook(output, start, end)
    print(f"Generated {output.resolve()}")


if __name__ == "__main__":
    main()
