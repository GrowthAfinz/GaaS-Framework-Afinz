#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera upsert SQL para serasa_marketplace_daily_funnel a partir do xlsx BI
(transposto: datas nas colunas; linhas Ofertas/Pedidos/Confirmados/Aprovados + %).

Uso: python gen_marketplace.py <xlsx> <saida.sql> <source_file_name> <lastClosedDate=YYYY-MM-DD>
"""
import sys, json, hashlib, datetime
import openpyxl

xlsx, out_sql, source_file, last_closed = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]

wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
rows = [list(r) for r in ws.iter_rows(values_only=True)]

# mapeia por rotulo da 1a celula
by_label = {}
for r in rows:
    key = (str(r[0]).strip() if r[0] is not None else "")
    by_label[key] = r[1:]

def parse_date(s):
    return datetime.datetime.strptime(str(s).strip(), "%b %d, %Y").date().isoformat()

dates = [parse_date(x) for x in by_label["Data"]]
def col(label): return by_label[label]

ofertas = col("Ofertas"); pedidos = col("Pedidos"); confirmados = col("Confirmados"); aprovados = col("Aprovados")
ped_pct = col("Pedidos %"); conf_pct = col("Confirmados %"); apr_pct = col("Aprovados %")

def as_int(v): return None if v in (None, "") else int(round(float(v)))
def as_rate(v):
    if v in (None, ""): return 0.0          # dia de abertura (denominador 0) -> rate 0 (col e NOT NULL)
    return round(float(v), 12)

records, problems = [], []
for i, d in enumerate(dates):
    o, p, c, a = as_int(ofertas[i]), as_int(pedidos[i]), as_int(confirmados[i]), as_int(aprovados[i])
    rec = dict(data=d, ofertas=o, pedidos=p, confirmados=c, aprovados=a,
               pedidosRateSource=as_rate(ped_pct[i]),
               confirmadosRateSource=as_rate(conf_pct[i]),
               aprovadosRateSource=as_rate(apr_pct[i]))
    # valida constraints da tabela
    if None in (o, p, c, a): problems.append(f"{d}: volume nulo {(o,p,c,a)}")
    elif not (o >= p >= c >= a >= 0): problems.append(f"{d}: funil nao-decrescente {(o,p,c,a)}")
    for rk in ("pedidosRateSource","confirmadosRateSource","aprovadosRateSource"):
        if not (0 <= rec[rk] <= 1): problems.append(f"{d}: {rk}={rec[rk]} fora de [0,1]")
    records.append(rec)

if problems:
    print("VIOLACOES DE CONSTRAINT (upsert falharia):")
    for pb in problems: print("  -", pb)
    raise SystemExit(1)

source_hash = hashlib.sha256(json.dumps(records).encode()).hexdigest()

def s(v): return "'" + str(v).replace("'", "''") + "'"
def n(v): return "null" if v is None else str(v)

vals = []
for r in records:
    partial = r["data"] > last_closed
    vals.append(f"""  ({s(r['data'])}, {n(r['ofertas'])}, {n(r['pedidos'])}, {n(r['pedidosRateSource'])}, {n(r['confirmados'])}, {n(r['confirmadosRateSource'])}, {n(r['aprovados'])}, {n(r['aprovadosRateSource'])}, {s('partial' if partial else 'complete')}, {"array['current_day_open']::text[]" if partial else "'{}'::text[]"}, {s(source_file)}, {s(source_hash)})""")

values_block = ",\n".join(vals)
sql = f"""insert into public.serasa_marketplace_daily_funnel (
  data, ofertas, pedidos, pedidos_rate_source, confirmados, confirmados_rate_source,
  aprovados, aprovados_rate_source, quality_status, quality_notes, source_file, source_hash
)
values
{values_block}
on conflict (data) do update set
  ofertas = excluded.ofertas,
  pedidos = excluded.pedidos,
  pedidos_rate_source = excluded.pedidos_rate_source,
  confirmados = excluded.confirmados,
  confirmados_rate_source = excluded.confirmados_rate_source,
  aprovados = excluded.aprovados,
  aprovados_rate_source = excluded.aprovados_rate_source,
  quality_status = excluded.quality_status,
  quality_notes = excluded.quality_notes,
  source_file = excluded.source_file,
  source_hash = excluded.source_hash,
  imported_at = now();
"""
open(out_sql, "w", encoding="utf-8").write(sql)
partials = [r["data"] for r in records if r["data"] > last_closed]
print(f"OK: {len(records)} dias ({records[0]['data']} .. {records[-1]['data']})")
print(f"  marcados 'partial' (>{last_closed}): {partials}")
print(f"  source_hash: {source_hash[:16]}...  source_file: {source_file}")
print(f"  SQL -> {out_sql}")
