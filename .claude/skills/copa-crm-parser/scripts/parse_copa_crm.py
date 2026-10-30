#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Copa CRM Parser — Afinz Growth Mídia
Processa arquivo de resultados SFMC (CSV ou XLSX multi-bloco)
e gera Copa2026_CRM.xlsx com abas Diário e Semanal.

Uso:
    python parse_copa_crm.py <arquivo_entrada> [arquivo_saida]

Entrada: CSV (sep=;) ou XLSX no formato multi-bloco Afinz
Saída:   Copa2026_CRM.xlsx (default: mesma pasta da entrada)
"""

import sys
import os
import re
from datetime import datetime, timedelta
from collections import defaultdict
from pathlib import Path

# ── dependências ──────────────────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Instale openpyxl: pip install openpyxl --break-system-packages")

# ══════════════════════════════════════════════════════════════════════════════
# TAXONOMIA — mapeamentos e regras com incongruências documentadas
# ══════════════════════════════════════════════════════════════════════════════

# Mapeamento de parceiro a partir de campos da activity name
# Incongruência 1: Bem Barato usa 'bb' (emails) e 'bbt' (WPP) — ambos válidos
# Incongruência 7: JOR_AQUISICAO_BB_ sem B2B2C — ainda é Bem Barato
PARTNER_FROM_ACTIVITY = {
    'bbt': 'Bem Barato',   # pos[2] em afz_car_bbt_aqs_wpp_...
    'bb':  'Bem Barato',   # pos[5] em afz_car_vis_aqs_email_bb_...
    'dia': 'DIA',
}

# Tokens que identificam parceiro nas journey names
PARTNER_FROM_JOURNEY = [
    ('PLURIX',    'Plurix'),      # Incongruência 2: usa nome completo, não 'plx'
    ('B2B2C_BB',  'Bem Barato'),
    ('B2B2C_DIA', 'DIA'),
    ('BB_',       'Bem Barato'),  # Incongruência 7: JOR_AQUISICAO_BB_ sem B2B2C
    ('_DIA_',     'DIA'),
    ('B2C',       'B2C'),
]

# Tokens que identificam tipo de comunicação
# Incongruência 8: journeys ABD (carrinho abandonado) usam _aqs_ no activity name
# mas são Reativação — identificadas pelo nome da journey (REATIVACAO ou ABD+PROPRI)
TIPO_TOKENS = [
    ('JOR_RENTABILIZACAO', 'Rentabilização'),
    ('_rnt_',              'Rentabilização'),
    ('sorteio',            'Rentabilização'),
    ('cartonista',         'Rentabilização'),
    ('_atv_sorteio',       'Rentabilização'),
    ('REATIVACAO',         'Reativação'),
    ('_rtv_',              'Reativação'),
    ('reativacao',         'Reativação'),
    # Carrinho abandonado: journey contém ABD — activity usa _aqs_ por engano
    # Padrão: JOR_AQUISICAO_B2C_*_ABD_REATIVACAO_* e JOR_AQUISICAO_B2C_PROPRI_ABD_*
    ('abd_reativacao',     'Reativação'),
    ('propri_abd',         'Reativação'),
    ('sab_abd',            'Reativação'),
    ('21d_abd',            'Reativação'),
]

# Finalizados ECRED — resultado real vem do GaaS (coluna Emitidos), não do SFMC PERF
# Chave: (activity_name, data_YYYY-MM-DD) → finalizados
# Incongruência 9: SFMC não captura resultado ECRED no bloco PERF para journeys ABD
ECRED_FINALIZADOS = {
    ('afz_car_vis_aqs_wpp_car_disp6copaecred_pontual',  '2026-04-29'): 43,
    ('afz_car_vis_aqs_wpp_car_disp7copaecred_pontual',  '2026-04-30'): 13,
    ('afz_car_vis_aqs_wpp_car_dispd22copaecred_pontual','2026-04-30'):  9,
    ('afz_car_vis_aqs_wpp_car_dispd22copaecred_pontual','2026-05-06'): 36,
    ('afz_car_vis_aqs_wpp_car_disp3copaecred_pontual',  '2026-05-13'): 35,
    ('afz_car_vis_aqs_wpp_car_disp6scopaecred_pontual', '2026-04-29'): 82,
    ('afz_car_vis_aqs_wpp_car_dispcopaecredsab_pontual','2026-05-09'): 51,
    ('afz_car_vis_aqs_wpp_car_dispcopaecredsab_pontual','2026-05-16'): 95,
}

# Mapa de público — pos[5] da activity name
# 'upo' não existe na taxonomia oficial (Incongruência 2)
# 'bb' e 'dia' aparecem como público em emails de parceiros (Incongruência 1)
PUBLICO_MAP = {
    'anc':  'Aprovados Não Convertidos',
    'ngd':  'Negados',
    'bsp':  'Base Proprietária',
    'bp':   'Base Proprietária',
    'apr':  'Aprovados',
    'car':  'Carrinho Abandonado',
    'atv':  'Ativo Geral',
    'rtv':  'Reativação',
    'rec':  'Recência',
    'upo':  'Upload',               # ad-hoc, não na taxonomia
    'bb':   'Base Bem Barato',      # público em emails BB
    'dia':  'Base DIA',             # público em emails DIA
    'lp':   'LP Opt-in',
    'srsa': 'Serasa',
    'sab':  'Sábado',               # ad-hoc recência
    'abd':  'Abandono',
    'emi':  'Clientes Emissores',
    'cart': 'Cartonistas',
    'club': 'Clube',
    'atl':  'Ativo com Limite',
    'freq': 'Frequentes e Recorrentes',
    'blq':  'Cartão Bloqueado',
    'rtv':  'Reativação',
}

# Segmentos ad-hoc — Incongruência 4
ADHOC_TOKENS = ['clus2', 'clustermar', 'ecred', 'copaint', 'copaecred',
                 '21d', 'sab', 'propri', 'fase1', 'diadasmaes', 'upo', 'srsa']

# Normalização de período — Incongruência 5
PERIODO_NORM = {
    'abri26': 'abr26', 'abr26': 'abr26',
    'maio26': 'mai26', 'mai26': 'mai26',
}

# Cores por parceiro (hex sem #)
PARTNER_COLORS = {
    'B2C':       'DBEAFE',
    'Bem Barato': 'DCFCE7',
    'Plurix':    'FEF3C7',
    'DIA':       'FFE4E6',
}


# ══════════════════════════════════════════════════════════════════════════════
# FUNÇÕES DE CLASSIFICAÇÃO
# ══════════════════════════════════════════════════════════════════════════════

def is_copa(val: str) -> bool:
    return bool(val and isinstance(val, str) and 'copa' in val.lower())


def classify_partner(activity: str, journey: str) -> str:
    act = str(activity or '').lower()
    jor = str(journey or '').upper()
    parts = act.split('_')

    # Empresa = plu → Plurix direto
    if parts and parts[0] == 'plu':
        return 'Plurix'

    # Testa pos[2] (parceiro na activity) e pos[5] (público que pode ser bb/dia)
    for pos in [2, 5]:
        if len(parts) > pos and parts[pos] in PARTNER_FROM_ACTIVITY:
            return PARTNER_FROM_ACTIVITY[parts[pos]]

    # Fallback: journey name
    for token, partner in PARTNER_FROM_JOURNEY:
        if token in jor:
            return partner

    return 'B2C'


def classify_tipo(activity: str, journey: str) -> str:
    combined = (str(activity or '') + str(journey or '')).lower()
    for token, tipo in TIPO_TOKENS:
        if token.lower() in combined:
            return tipo
    return 'Aquisição'


def detect_adhoc(activity: str) -> str:
    act = activity.lower()
    found = [t for t in ADHOC_TOKENS if t in act]
    return '|'.join(found) if found else ''


def extract_publico(activity: str) -> str:
    """Extrai público da posição 5 da activity name e traduz pelo dicionário."""
    parts = str(activity or '').lower().split('_')
    raw = parts[5] if len(parts) > 5 else ''
    return PUBLICO_MAP.get(raw, raw) if raw else ''


def extract_nome_jornada(journey: str) -> str:
    """Retorna nome legível da jornada a partir do journey name SFMC.
    Remove prefixo JOR_ e substitui _ por espaço, title case."""
    j = str(journey or '').strip()
    if j.upper().startswith('JOR_'):
        j = j[4:]
    return j.replace('_', ' ').title()


def parse_date(val) -> datetime | None:
    if not val:
        return None
    if isinstance(val, str):
        val = val.strip()
        for fmt in ['%d/%m/%Y', '%Y-%m-%d', '%d/%m/%y']:
            try:
                return datetime.strptime(val, fmt)
            except ValueError:
                pass
    if hasattr(val, 'year'):
        return val
    return None


def week_start(dt: datetime) -> str:
    return (dt - timedelta(days=dt.weekday())).strftime('%Y-%m-%d')


def safe_num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        v = v.strip().replace(',', '.')
        try:
            return float(v)
        except ValueError:
            pass
    return None


# ══════════════════════════════════════════════════════════════════════════════
# LEITURA DO ARQUIVO DE RESULTADOS
# ══════════════════════════════════════════════════════════════════════════════

def load_rows(filepath: str):
    """Lê arquivo CSV ou XLSX e retorna lista de tuplas (row) com valores brutos."""
    ext = Path(filepath).suffix.lower()

    if ext == '.xlsx':
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=4, values_only=True))
        wb.close()
        return rows

    elif ext == '.csv':
        import csv
        rows = []
        with open(filepath, encoding='utf-8-sig', newline='') as f:
            # Detecta separador
            sample = f.read(2048)
            f.seek(0)
            sep = ';' if sample.count(';') > sample.count(',') else ','
            reader = csv.reader(f, delimiter=sep)
            all_rows = list(reader)
        # Pula primeiras 3 linhas (header/pivot), começa na linha 4
        return [tuple(r) for r in all_rows[3:]]

    else:
        sys.exit(f"Formato não suportado: {ext}. Use .xlsx ou .csv")


def safe_get(row, idx):
    try:
        return row[idx]
    except IndexError:
        return None


# ══════════════════════════════════════════════════════════════════════════════
# EXTRAÇÃO DOS BLOCOS
# ══════════════════════════════════════════════════════════════════════════════

def extract_blocks(rows):
    """
    Layout multi-bloco (1-based → 0-indexed):
    WPP:   0,1,2,3,4,5
    EMAIL: 8,9,10,11,12,13,14
    SMS:   16,17,18,19,20,21
    PERF:  23,24,25,26,27,28,29
    PUSH:  35,36,37,38,39
    """
    crm_rows = {}   # (activity, date_str) → dict CRM
    perf_rows = {}  # (activity, date_str) → dict PERF

    for row in rows:
        # ── WPP ──────────────────────────────────────────────────────────
        j = safe_get(row, 0)
        if is_copa(j):
            sent = safe_num(safe_get(row, 3)) or 0
            if sent > 0:
                act = str(safe_get(row, 1) or '').strip()
                dt  = parse_date(safe_get(row, 2))
                if act and dt:
                    key = (act, dt.strftime('%Y-%m-%d'))
                    crm_rows[key] = {
                        'canal': 'WPP', 'journey': j, 'activity': act, 'data': dt,
                        'partner': classify_partner(act, j),
                        'tipo': classify_tipo(act, j),
                        'publico': extract_publico(act),
                        'nome_jornada': extract_nome_jornada(j),
                        'adhoc': detect_adhoc(act),
                        'enviados':      int(sent),
                        'entregues':     int(safe_num(safe_get(row, 4)) or 0),
                        'abertos_lidos': int(safe_num(safe_get(row, 5)) or 0),
                        'cliques':       None,
                    }

        # ── EMAIL ─────────────────────────────────────────────────────────
        j = safe_get(row, 8)
        if is_copa(j):
            sent = safe_num(safe_get(row, 11)) or 0
            if sent > 0:
                act = str(safe_get(row, 9) or '').strip()
                dt  = parse_date(safe_get(row, 10))
                if act and dt:
                    key = (act, dt.strftime('%Y-%m-%d'))
                    crm_rows[key] = {
                        'canal': 'EMAIL', 'journey': j, 'activity': act, 'data': dt,
                        'partner': classify_partner(act, j),
                        'tipo': classify_tipo(act, j),
                        'publico': extract_publico(act),
                        'nome_jornada': extract_nome_jornada(j),
                        'adhoc': detect_adhoc(act),
                        'enviados':      int(sent),
                        'entregues':     int(safe_num(safe_get(row, 12)) or 0),
                        'abertos_lidos': int(safe_num(safe_get(row, 13)) or 0),
                        'cliques':       int(safe_num(safe_get(row, 14)) or 0),
                    }

        # ── SMS ───────────────────────────────────────────────────────────
        j = safe_get(row, 16)
        if is_copa(j):
            sent = safe_num(safe_get(row, 19)) or 0
            if sent > 0:
                act = str(safe_get(row, 17) or '').strip()
                dt  = parse_date(safe_get(row, 18))
                if act and dt:
                    key = (act, dt.strftime('%Y-%m-%d'))
                    crm_rows[key] = {
                        'canal': 'SMS', 'journey': j, 'activity': act, 'data': dt,
                        'partner': classify_partner(act, j),
                        'tipo': classify_tipo(act, j),
                        'publico': extract_publico(act),
                        'nome_jornada': extract_nome_jornada(j),
                        'adhoc': detect_adhoc(act),
                        'enviados':      int(sent),
                        'entregues':     int(safe_num(safe_get(row, 20)) or 0),
                        'abertos_lidos': None,
                        'cliques':       None,
                    }

        # ── PERF ──────────────────────────────────────────────────────────
        j = safe_get(row, 23)
        if is_copa(j):
            act = str(safe_get(row, 24) or '').strip()
            dt  = parse_date(safe_get(row, 25))
            if act and dt:
                key = (act, dt.strftime('%Y-%m-%d'))
                perf_rows[key] = {
                    'canal_perf':  str(safe_get(row, 26) or '').strip(),
                    'journey_perf': j,
                    'proposta':    int(safe_num(safe_get(row, 27)) or 0),
                    'aprovados':   int(safe_num(safe_get(row, 28)) or 0),
                    'finalizados': int(safe_num(safe_get(row, 29)) or 0),
                }

        # ── PUSH ──────────────────────────────────────────────────────────
        j = safe_get(row, 35)
        if is_copa(j):
            sent = safe_num(safe_get(row, 38)) or 0
            if sent > 0:
                act = str(safe_get(row, 36) or '').strip()
                dt  = parse_date(safe_get(row, 37))
                if act and dt:
                    key = (act, dt.strftime('%Y-%m-%d'))
                    crm_rows[key] = {
                        'canal': 'PUSH', 'journey': j, 'activity': act, 'data': dt,
                        'partner': classify_partner(act, j),
                        'tipo': classify_tipo(act, j),
                        'adhoc': detect_adhoc(act),
                        'enviados':      int(sent),
                        'entregues':     int(safe_num(safe_get(row, 39)) or 0),
                        'abertos_lidos': None,
                        'cliques':       None,
                    }

    return crm_rows, perf_rows


# ══════════════════════════════════════════════════════════════════════════════
# JOIN
# ══════════════════════════════════════════════════════════════════════════════

def join_blocks(crm_rows, perf_rows):
    joined = []
    all_keys = set(crm_rows) | set(perf_rows)

    for key in all_keys:
        c = crm_rows.get(key, {})
        p = perf_rows.get(key, {})

        act = key[0]
        dt  = c.get('data') or parse_date(key[1])
        j   = c.get('journey') or p.get('journey_perf', '')

        partner = c.get('partner') or classify_partner(act, j)
        tipo    = c.get('tipo')    or classify_tipo(act, j)
        canal   = c.get('canal')   or p.get('canal_perf', 'PERF')

        # Finalizados: PERF block + ECRED manual (GaaS) se existir
        data_str = dt.strftime('%Y-%m-%d') if dt else key[1]
        fin_perf = p.get('finalizados') or 0
        fin_ecred = ECRED_FINALIZADOS.get((act, data_str), 0)
        finalizados = (fin_perf + fin_ecred) if (fin_perf + fin_ecred) > 0 else None

        joined.append({
            'data':          data_str,
            'semana_inicio': week_start(dt) if dt else '',
            'partner':       partner,
            'tipo':          tipo,
            'canal':         canal,
            'publico':       c.get('publico', extract_publico(act)),
            'nome_jornada':  c.get('nome_jornada', extract_nome_jornada(j)),
            'journey':       j,
            'activity':      act,
            'adhoc':         c.get('adhoc', detect_adhoc(act)),
            'enviados':      c.get('enviados'),
            'entregues':     c.get('entregues'),
            'abertos_lidos': c.get('abertos_lidos'),
            'cliques':       c.get('cliques'),
            'proposta':      p.get('proposta'),
            'aprovados':     p.get('aprovados'),
            'finalizados':   finalizados,
            '_dt':           dt,
        })

    joined.sort(key=lambda r: r['_dt'] or datetime.min)
    return joined


# ══════════════════════════════════════════════════════════════════════════════
# ESTILOS EXCEL
# ══════════════════════════════════════════════════════════════════════════════

HDR_FILL  = PatternFill('solid', start_color='0033A0')
HDR_FONT  = Font(bold=True, color='FFFFFF', name='Arial', size=10)
BODY_FONT = Font(name='Arial', size=10)
ALT_FILL  = PatternFill('solid', start_color='EEF2FF')
TOT_FILL  = PatternFill('solid', start_color='FFF9C4')
TOT_FONT  = Font(bold=True, name='Arial', size=10)
CENTER    = Alignment(horizontal='center', vertical='center')
thin      = Side(style='thin', color='D0D7E5')
BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)


def apply_header(ws, headers, row=1):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BORDER


def style_row(ws, r, ncols, alt=False, total=False):
    fill = TOT_FILL if total else (ALT_FILL if alt else None)
    font = TOT_FONT if total else BODY_FONT
    for col in range(1, ncols + 1):
        c = ws.cell(row=r, column=col)
        c.font = font
        c.border = BORDER
        if fill:
            c.fill = fill


def color_partner_cell(ws, row, col, partner):
    color = PARTNER_COLORS.get(partner)
    if color:
        ws.cell(row=row, column=col).fill = PatternFill('solid', start_color=color)


# ══════════════════════════════════════════════════════════════════════════════
# GERAÇÃO DO XLSX
# ══════════════════════════════════════════════════════════════════════════════

DAILY_HEADERS = [
    'Data', 'Semana', 'Partner', 'Tipo', 'Canal', 'Público', 'Nome Jornada',
    'Activity', 'Segmento Adhoc',
    'Enviados', 'Entregues', 'Abertos/Lidos', 'Cliques',
    'Proposta', 'Aprovados', 'Finalizados',
    'Tx Entrega', 'Tx Abertura', 'Tx Conv Fin',
]

WEEKLY_HEADERS = [
    'Semana Início', 'Partner', 'Tipo', 'Canal', 'Qtd Disparos',
    'Enviados', 'Entregues', 'Abertos/Lidos', 'Cliques',
    'Proposta', 'Aprovados', 'Finalizados',
    'Tx Entrega', 'Tx Abertura', 'Tx Conv Fin',
]

# Diário: 19 colunas | Semanal: 15 colunas (sem Público/NomeJornada — agrega diferente)
DAILY_COL_WIDTHS  = [12,12,13,16,8,22,36, 50,18, 10,10,13,9, 10,10,12, 10,11,11]
WEEKLY_COL_WIDTHS = [14,13,16,8,10, 10,10,13,9, 10,10,12, 10,11,11]


def build_xlsx(joined, output_path):
    wb = openpyxl.Workbook()

    # ── ABA DIÁRIO ────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'Diário'
    apply_header(ws1, DAILY_HEADERS)
    ws1.freeze_panes = 'A2'
    ws1.auto_filter.ref = f'A1:{get_column_letter(len(DAILY_HEADERS))}1'

    for i, r in enumerate(joined, 2):
        env  = r['enviados']
        ent  = r['entregues']
        abi  = r['abertos_lidos']
        cli  = r['cliques']
        prop = r['proposta']
        fin  = r['finalizados']

        # Colunas após inserção de Público(F) e Nome Jornada(G):
        # Data=A Semana=B Partner=C Tipo=D Canal=E Público=F NomeJornada=G
        # Activity=H Adhoc=I Enviados=J Entregues=K Abertos=L Cliques=M
        # Proposta=N Aprovados=O Finalizados=P TxEnt=Q TxAbi=R TxConv=S
        tx_ent = f'=IF(AND(J{i}>0,K{i}<>""),K{i}/J{i},"")'
        tx_abi = f'=IF(AND(K{i}>0,L{i}<>""),L{i}/K{i},"")'
        tx_con = f'=IF(AND(N{i}>0,P{i}<>""),P{i}/N{i},"")'

        row_vals = [
            r['data'], r['semana_inicio'], r['partner'], r['tipo'], r['canal'],
            r.get('publico', ''), r.get('nome_jornada', ''),
            r['activity'], r['adhoc'] or '',
            env, ent, abi, cli, prop, r.get('aprovados'), fin,
            tx_ent if env else '',
            tx_abi if (ent and abi is not None) else '',
            tx_con if (prop and fin is not None) else '',
        ]

        for col, v in enumerate(row_vals, 1):
            ws1.cell(row=i, column=col, value=v)

        style_row(ws1, i, len(DAILY_HEADERS), alt=(i % 2 == 0))
        color_partner_cell(ws1, i, 3, r['partner'])

        for col in [17, 18, 19]:
            ws1.cell(row=i, column=col).number_format = '0.0%'
        for col in [10, 11, 12, 13, 14, 15, 16]:
            c = ws1.cell(row=i, column=col)
            c.number_format = '#,##0'
            c.alignment = CENTER

    for i, w in enumerate(DAILY_COL_WIDTHS, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ── ABA SEMANAL ───────────────────────────────────────────────────────
    ws2 = wb.create_sheet('Semanal')
    apply_header(ws2, WEEKLY_HEADERS)
    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = f'A1:{get_column_letter(len(WEEKLY_HEADERS))}1'

    weekly = defaultdict(lambda: {
        'enviados': 0, 'entregues': 0, 'abertos_lidos': 0, 'cliques': 0,
        'proposta': 0, 'aprovados': 0, 'finalizados': 0, 'n': 0,
    })
    for r in joined:
        key = (r['semana_inicio'], r['partner'], r['tipo'], r['canal'])
        d = weekly[key]
        d['n'] += 1
        for f in ['enviados','entregues','abertos_lidos','cliques','proposta','aprovados','finalizados']:
            v = r.get(f)
            if isinstance(v, (int, float)):
                d[f] += v

    weekly_sorted = sorted(weekly.items(), key=lambda x: x[0])

    for i, (key, d) in enumerate(weekly_sorted, 2):
        semana, partner, tipo, canal = key
        env  = d['enviados']  or None
        ent  = d['entregues'] or None
        abi  = d['abertos_lidos'] or None
        cli  = d['cliques']   or None
        prop = d['proposta']  or None
        apr  = d['aprovados'] or None
        fin  = d['finalizados'] or None

        tx_ent = f'=IF(AND(F{i}>0,G{i}<>""),G{i}/F{i},"")'
        tx_abi = f'=IF(AND(G{i}>0,H{i}<>""),H{i}/G{i},"")'
        tx_con = f'=IF(AND(J{i}>0,L{i}<>""),L{i}/J{i},"")'

        row_vals = [semana, partner, tipo, canal, d['n'],
                    env, ent, abi, cli, prop, apr, fin,
                    tx_ent if env else '',
                    tx_abi if (ent and abi) else '',
                    tx_con if (prop and fin) else '']

        for col, v in enumerate(row_vals, 1):
            ws2.cell(row=i, column=col, value=v)

        style_row(ws2, i, len(WEEKLY_HEADERS), alt=(i % 2 == 0))
        color_partner_cell(ws2, i, 2, partner)

        for col in [13, 14, 15]:
            ws2.cell(row=i, column=col).number_format = '0.0%'
        for col in [6, 7, 8, 9, 10, 11, 12]:
            c = ws2.cell(row=i, column=col)
            c.number_format = '#,##0'
            c.alignment = CENTER

    # Linha de totais
    tr = len(weekly_sorted) + 2
    ws2.cell(tr, 1, 'TOTAL').font = TOT_FONT
    col_letters = {6:'F', 7:'G', 8:'H', 9:'I', 10:'J', 11:'K', 12:'L'}
    for col, letter in col_letters.items():
        c = ws2.cell(tr, col, f'=SUM({letter}2:{letter}{tr-1})')
        c.number_format = '#,##0'
        c.alignment = CENTER
    ws2.cell(tr, 13, f'=IF(F{tr}>0,G{tr}/F{tr},"")').number_format = '0.0%'
    ws2.cell(tr, 15, f'=IF(J{tr}>0,L{tr}/J{tr},"")').number_format = '0.0%'
    style_row(ws2, tr, len(WEEKLY_HEADERS), total=True)

    for i, w in enumerate(WEEKLY_COL_WIDTHS, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(output_path)
    return len(joined), len(weekly_sorted)




# ══════════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    input_path = sys.argv[1]
    if not os.path.exists(input_path):
        sys.exit(f"Arquivo nao encontrado: {input_path}")

    # Output path
    if len(sys.argv) >= 3:
        output_path = sys.argv[2]
    else:
        out_dir = os.path.dirname(os.path.abspath(input_path))
        output_path = os.path.join(out_dir, 'Copa2026_CRM.xlsx')

    print(f"Lendo: {input_path}")
    rows = load_rows(input_path)
    print(f"  Linhas brutas: {len(rows)}")

    crm_rows, perf_rows = extract_blocks(rows)
    print(f"  CRM unicos (activity+data): {len(crm_rows)}")
    print(f"  PERF unicos (activity+data): {len(perf_rows)}")

    matched  = len(set(crm_rows) & set(perf_rows))
    crm_only = len(set(crm_rows) - set(perf_rows))
    print(f"  Join completo (CRM+PERF): {matched}")
    print(f"  So CRM (rentabilizacao sem funil): {crm_only}")

    joined = join_blocks(crm_rows, perf_rows)

    # Resumo de parceiros e periodo
    partners = set(r['partner'] for r in joined)
    datas = [r['_dt'] for r in joined if r.get('_dt')]
    periodo = f"{min(datas).strftime('%d/%m/%Y')} -> {max(datas).strftime('%d/%m/%Y')}" if datas else 'N/A'
    total_env = sum(r['enviados'] or 0 for r in joined)
    total_fin = sum(r['finalizados'] or 0 for r in joined)

    print(f"  Parceiros: {', '.join(sorted(partners))}")
    print(f"  Periodo: {periodo}")
    print(f"  Total enviados: {total_env:,}")
    print(f"  Total finalizados: {total_fin:,}")

    print(f"Gerando: {output_path}")
    n_daily, n_weekly = build_xlsx(joined, output_path)
    print(f"  Aba Diario: {n_daily} linhas")
    print(f"  Aba Semanal: {n_weekly} linhas + total")
    print(f"Arquivo salvo: {output_path}")

    import json
    result = {
        'output': output_path,
        'linhas_diario': n_daily,
        'linhas_semanal': n_weekly,
        'parceiros': sorted(partners),
        'periodo': periodo,
        'total_enviados': total_env,
        'total_finalizados': total_fin,
        'join_completo': matched,
        'crm_only': crm_only,
    }
    print(f"JSON: {json.dumps(result, ensure_ascii=False)}")
    return result


if __name__ == '__main__':
    main()
